import json

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.db import models, transaction
from django.db.models import Sum
from django.shortcuts import get_object_or_404, redirect, render

from catalogue.models import Produit
from stock.models import SoldeStock

from core.models import Profil
from .models import StatutTransfert, Transfert, TransfertLigne, TypeTransfert
from .services import accepter_transfert, destinations_possibles, envoyer_transfert, rejeter_transfert


def _filtrer_transferts(request):
    """Retourne (résultat, filtres) en appliquant tous les filtres GET."""
    sites = request.user.sites_autorises()
    site_ids = set(sites.values_list("pk", flat=True))

    sens_filtre    = request.GET.get("sens", "")
    statut_filtre  = request.GET.get("statut", "")
    debut          = request.GET.get("debut", "")
    fin            = request.GET.get("fin", "")
    origine_filtre = request.GET.get("origine", "")
    dest_filtre    = request.GET.get("destination", "")

    qs = (
        Transfert.objects
        .filter(
            models.Q(site_origine__in=sites) |
            models.Q(site_destination__in=sites) |
            models.Q(type_transfert__in=[TypeTransfert.DON, TypeTransfert.SURPLUS], site_origine__in=sites)
        )
        .select_related("site_origine", "site_destination", "cree_par")
        .order_by("-cree_le")
        .distinct()
    )
    if statut_filtre:
        qs = qs.filter(statut=statut_filtre)
    if debut:
        qs = qs.filter(cree_le__date__gte=debut)
    if fin:
        qs = qs.filter(cree_le__date__lte=fin)
    if origine_filtre:
        qs = qs.filter(site_origine_id=origine_filtre)
    if dest_filtre:
        qs = qs.filter(site_destination_id=dest_filtre)

    result = []
    for t in qs:
        t.sens = "envoye" if t.site_origine_id in site_ids else "recu"
        if sens_filtre and t.sens != sens_filtre:
            continue
        result.append(t)

    filtres = {
        "sens": sens_filtre, "statut": statut_filtre,
        "debut": debut, "fin": fin,
        "origine": origine_filtre, "destination": dest_filtre,
    }
    return result, filtres


@login_required
def transferts_liste(request):
    from core.models import Profil
    transferts, filtres = _filtrer_transferts(request)
    sites_filtre = (
        request.user.sites_autorises().order_by("type", "nom")
        if request.user.profil == Profil.DG or request.user.is_superuser
        else None
    )
    return render(request, "transferts/liste.html", {
        "transferts": transferts[:100],
        "statuts": StatutTransfert.choices,
        "sites_filtre": sites_filtre,
        "filtres": filtres,
    })


@login_required
def export_transferts(request):
    import io
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from django.http import HttpResponse

    transferts, _ = _filtrer_transferts(request)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transferts"

    entetes = ["#", "Sens", "Expéditeur", "Destinataire", "Statut", "Date", "Par", "Traité le", "Par"]
    for col, titre in enumerate(entetes, 1):
        c = ws.cell(row=1, column=col, value=titre)
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor="16233F")
        c.alignment = Alignment(horizontal="center", vertical="center")

    for row, t in enumerate(transferts, 2):
        ws.cell(row=row, column=1, value=f"TRF-{t.pk}")
        ws.cell(row=row, column=2, value="Envoyé" if t.sens == "envoye" else "Reçu")
        ws.cell(row=row, column=3, value=t.site_origine.nom)
        ws.cell(row=row, column=4, value=t.site_destination.nom)
        ws.cell(row=row, column=5, value=t.get_statut_display())
        ws.cell(row=row, column=6, value=t.cree_le.strftime("%d/%m/%Y %H:%M") if t.cree_le else "")
        ws.cell(row=row, column=7, value=t.cree_par.get_full_name() or t.cree_par.username)
        ws.cell(row=row, column=8, value=t.traite_le.strftime("%d/%m/%Y %H:%M") if t.traite_le else "")
        ws.cell(row=row, column=9, value=t.traite_par.get_full_name() if t.traite_par else "")

    for col in ws.columns:
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(
            max(len(str(c.value or "")) for c in col) + 4, 50
        )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(buf, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="transferts.xlsx"'
    return response


@login_required
def transfert_formulaire(request):
    if not request.user.site_id:
        messages.error(request, "Votre compte n'est pas rattaché à un site — impossible de créer un transfert.")
        return redirect("transferts_liste")

    site_origine = request.user.site
    destinations = destinations_possibles(request.user).order_by("type", "nom")
    produits = Produit.objects.filter(actif=True).order_by("code")

    soldes = {
        s.produit_id: s.quantite
        for s in SoldeStock.objects.filter(site=site_origine)
    }
    stocks_disponibles = dict(soldes)

    if request.method == "POST":
        type_t = request.POST.get("type_transfert", TypeTransfert.NORMAL)
        if type_t not in TypeTransfert.values:
            type_t = TypeTransfert.NORMAL
        destination_id = request.POST.get("site_destination") if type_t == TypeTransfert.NORMAL else None
        observations = request.POST.get("observations", "").strip()
        produit_ids = request.POST.getlist("produit_id")
        quantites = request.POST.getlist("quantite")

        lignes = []
        erreurs = []
        for pid, q in zip(produit_ids, quantites):
            try:
                q_int = int(q)
                pid_int = int(pid)
                if q_int > 0 and pid:
                    dispo = stocks_disponibles.get(pid_int, 0)
                    if q_int > dispo:
                        erreurs.append(f"Quantité ({q_int}) supérieure au stock disponible ({dispo}) pour le produit #{pid_int}.")
                    else:
                        lignes.append((pid_int, q_int))
            except (ValueError, TypeError):
                pass

        destination_valide = (
            type_t != TypeTransfert.NORMAL or
            destinations.filter(pk=destination_id).exists()
        )

        if erreurs:
            for e in erreurs:
                messages.error(request, e)
        elif not lignes:
            messages.error(request, "Ajoutez au moins une ligne.")
        elif not destination_valide:
            messages.error(request, "Destination non autorisée.")
        else:
            try:
                with transaction.atomic():
                    t = Transfert.objects.create(
                        site_origine=site_origine,
                        site_destination_id=destination_id,
                        type_transfert=type_t,
                        observations=observations,
                        cree_par=request.user,
                    )
                    TransfertLigne.objects.bulk_create([
                        TransfertLigne(transfert=t, produit_id=pid, quantite=q)
                        for pid, q in lignes
                    ])
                    envoyer_transfert(t, par=request.user)
                if type_t == TypeTransfert.NORMAL:
                    messages.success(request, "Transfert envoyé — stock débité, le destinataire a été notifié.")
                else:
                    label = "Don" if type_t == TypeTransfert.DON else "Surplus"
                    messages.success(request, f"Transfert {label} créé — en attente de validation par le DG.")
                return redirect("transfert_detail", pk=t.pk)
            except ValidationError as e:
                messages.error(request, str(e))
            except Exception as e:
                messages.error(request, str(e))

    return render(request, "transferts/formulaire.html", {
        "site_origine": site_origine,
        "destinations": destinations,
        "produits": produits,
        "stocks_json": json.dumps(stocks_disponibles),
        "types_transfert": TypeTransfert.choices,
    })


@login_required
def transfert_detail(request, pk):
    sites = request.user.sites_autorises()
    t = get_object_or_404(
        Transfert.objects.select_related("site_origine", "site_destination", "cree_par", "traite_par"),
        pk=pk,
    )
    site_ids_impliques = [t.site_origine_id]
    if t.site_destination_id:
        site_ids_impliques.append(t.site_destination_id)
    est_dg = request.user.profil == Profil.DG or request.user.is_superuser
    if not sites.filter(pk__in=site_ids_impliques).exists() and not est_dg:
        messages.error(request, "Accès refusé.")
        return redirect("transferts_liste")

    if t.est_special:
        peut_traiter = est_dg and t.statut == StatutTransfert.EN_ATTENTE
    else:
        est_destinataire = sites.filter(pk=t.site_destination_id).exists()
        peut_traiter = est_destinataire and t.statut == StatutTransfert.EN_ATTENTE

    lignes = t.lignes.select_related("produit").order_by("produit__code")
    return render(request, "transferts/detail.html", {
        "transfert": t,
        "lignes": lignes,
        "peut_traiter": peut_traiter,
    })


@login_required
def transfert_accepter(request, pk):
    t = get_object_or_404(Transfert, pk=pk, statut=StatutTransfert.EN_ATTENTE)
    est_dg = request.user.profil == Profil.DG or request.user.is_superuser
    if t.est_special:
        if not est_dg:
            messages.error(request, "Seul le DG peut valider un transfert Don ou Surplus.")
            return redirect("transfert_detail", pk=pk)
    else:
        if not request.user.sites_autorises().filter(pk=t.site_destination_id).exists():
            messages.error(request, "Vous ne pouvez accepter que les transferts destinés à votre site.")
            return redirect("transfert_detail", pk=pk)
    if request.method == "POST":
        try:
            accepter_transfert(t, par=request.user)
            if t.est_special:
                messages.success(request, f"Transfert {t.get_type_transfert_display()} validé.")
            else:
                messages.success(request, "Transfert accepté — stock crédité.")
        except ValidationError as e:
            messages.error(request, str(e))
    return redirect("transfert_detail", pk=pk)


@login_required
def transfert_rejeter(request, pk):
    t = get_object_or_404(Transfert, pk=pk, statut=StatutTransfert.EN_ATTENTE)
    est_dg = request.user.profil == Profil.DG or request.user.is_superuser
    if t.est_special:
        if not est_dg:
            messages.error(request, "Seul le DG peut rejeter un transfert Don ou Surplus.")
            return redirect("transfert_detail", pk=pk)
    elif not request.user.sites_autorises().filter(pk=t.site_destination_id).exists():
        messages.error(request, "Vous ne pouvez rejeter que les transferts destinés à votre site.")
        return redirect("transfert_detail", pk=pk)
    if request.method == "POST":
        try:
            rejeter_transfert(t, par=request.user)
            messages.success(request, "Transfert rejeté — stock recrédité à l'émetteur.")
        except ValidationError as e:
            messages.error(request, str(e))
    return redirect("transfert_detail", pk=pk)
