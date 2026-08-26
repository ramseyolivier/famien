"""Exports Excel pour les achats fournisseurs."""

import io

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.utils import timezone

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import CommandeFournisseur, Reception


def _style_entete(cell):
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.fill = PatternFill("solid", fgColor="16233F")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _ajuster_colonnes(ws):
    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 50)


@login_required
def export_commande_detail_excel(request, pk):
    from .models import CommandeFournisseur

    cf = CommandeFournisseur.objects.select_related(
        "fournisseur", "site_destination", "cree_par"
    ).prefetch_related("lignes__produit").get(pk=pk)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Commande CF #{cf.pk}"
    ws.freeze_panes = "A2"

    entetes = ["Produit", "Quantité commandée"]
    for col, titre in enumerate(entetes, 1):
        cell = ws.cell(row=1, column=col, value=titre)
        _style_entete(cell)
    ws.row_dimensions[1].height = 30

    row = 2
    total = 0
    for ligne in cf.lignes.select_related("produit").order_by("produit__code"):
        ws.cell(row=row, column=1, value=f"{ligne.produit.code} — {ligne.produit.designation}")
        ws.cell(row=row, column=2, value=ligne.quantite)
        total += ligne.quantite
        row += 1

    ws.cell(row=row, column=1, value="Total").font = Font(bold=True)
    cell_total = ws.cell(row=row, column=2, value=total)
    cell_total.font = Font(bold=True)

    _ajuster_colonnes(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nom = f"commande_{cf.pk}_{timezone.localdate().isoformat()}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nom}"'
    return response


@login_required
def export_commandes_fournisseur_excel(request):
    from core.models import Profil
    from django.db.models import Q

    u = request.user
    sites = u.sites_autorises()

    qs = (
        CommandeFournisseur.objects
        .filter(site_destination__in=sites)
        .filter(Q(statut="BROUILLON", cree_par=u) | ~Q(statut="BROUILLON"))
        .select_related("fournisseur", "site_destination", "cree_par")
        .prefetch_related("lignes__produit")
        .order_by("-cree_le")
    )

    if request.GET.get("commande", "").isdigit():
        qs = qs.filter(pk=request.GET["commande"])
    if request.GET.get("statut"):
        qs = qs.filter(statut=request.GET["statut"])
    if request.GET.get("debut"):
        qs = qs.filter(cree_le__date__gte=request.GET["debut"])
    if request.GET.get("fin"):
        qs = qs.filter(cree_le__date__lte=request.GET["fin"])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Commandes fournisseur"
    ws.freeze_panes = "A2"

    entetes = ["#", "Fournisseur", "Site destination", "Statut", "Date", "Créé par", "Produit", "Code", "Qté commandée"]
    for col, titre in enumerate(entetes, 1):
        cell = ws.cell(row=1, column=col, value=titre)
        _style_entete(cell)
    ws.row_dimensions[1].height = 30

    row = 2
    for commande in qs:
        lignes = list(commande.lignes.all())
        base = [
            commande.pk,
            commande.fournisseur.raison_sociale,
            commande.site_destination.nom,
            commande.get_statut_display(),
            commande.cree_le.strftime("%d/%m/%Y"),
            commande.cree_par.get_full_name() or commande.cree_par.username,
        ]
        if not lignes:
            for col, val in enumerate(base, 1):
                ws.cell(row=row, column=col, value=val)
            row += 1
        else:
            for ligne in lignes:
                for col, val in enumerate(base, 1):
                    ws.cell(row=row, column=col, value=val)
                ws.cell(row=row, column=7, value=ligne.produit.designation)
                ws.cell(row=row, column=8, value=ligne.produit.code)
                ws.cell(row=row, column=9, value=ligne.quantite)
                row += 1

    _ajuster_colonnes(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nom = f"commandes_fournisseur_{timezone.localdate().isoformat()}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nom}"'
    return response


@login_required
def export_reception_detail_excel(request, pk):
    from .models import Reception

    rec = Reception.objects.select_related(
        "commande__fournisseur", "commande__site_destination",
        "fournisseur", "site_destination", "cree_par", "valide_par",
    ).prefetch_related("lignes__produit").get(pk=pk)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Réception {rec.titre}"
    ws.freeze_panes = "A2"

    if rec.commande_id:
        entetes = ["Produit", "Commandé", "Reçu", "Montant (F CFA)", "Écart", "Justification"]
    else:
        entetes = ["Produit", "Reçu", "Montant (F CFA)"]

    for col, titre in enumerate(entetes, 1):
        cell = ws.cell(row=1, column=col, value=titre)
        _style_entete(cell)
    ws.row_dimensions[1].height = 30

    row = 2
    total_montant = 0
    for ligne in rec.lignes.select_related("produit").order_by("produit__code"):
        montant = float(ligne.montant)
        total_montant += montant
        produit_label = f"{ligne.produit.code} — {ligne.produit.designation}"
        if rec.commande_id:
            ecart = ligne.quantite_recue - ligne.quantite_attendue
            ws.cell(row=row, column=1, value=produit_label)
            ws.cell(row=row, column=2, value=ligne.quantite_attendue)
            ws.cell(row=row, column=3, value=ligne.quantite_recue)
            ws.cell(row=row, column=4, value=montant)
            ws.cell(row=row, column=5, value=ecart if ecart != 0 else "—")
            ws.cell(row=row, column=6, value=ligne.justification_ecart or "—")
        else:
            ws.cell(row=row, column=1, value=produit_label)
            ws.cell(row=row, column=2, value=ligne.quantite_recue)
            ws.cell(row=row, column=3, value=montant)
        row += 1

    # Ligne total
    col_montant = 4 if rec.commande_id else 3
    ws.cell(row=row, column=col_montant - 1, value="Total").font = Font(bold=True)
    cell_total = ws.cell(row=row, column=col_montant, value=total_montant)
    cell_total.font = Font(bold=True)

    _ajuster_colonnes(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nom = f"reception_{rec.pk}_{timezone.localdate().isoformat()}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nom}"'
    return response


@login_required
def export_receptions_fournisseur_excel(request):
    from django.db.models import Q

    u = request.user
    sites = u.sites_autorises()

    qs = (
        Reception.objects
        .select_related("commande__fournisseur", "commande__site_destination",
                        "fournisseur", "site_destination", "cree_par")
        .prefetch_related("lignes__produit")
        .order_by("-cree_le")
    )

    if request.GET.get("reception", "").isdigit():
        qs = qs.filter(pk=request.GET["reception"])
    if request.GET.get("statut"):
        qs = qs.filter(statut=request.GET["statut"])
    if request.GET.get("debut"):
        qs = qs.filter(cree_le__date__gte=request.GET["debut"])
    if request.GET.get("fin"):
        qs = qs.filter(cree_le__date__lte=request.GET["fin"])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Réceptions fournisseur"
    ws.freeze_panes = "A2"

    entetes = ["#", "Commande", "Fournisseur", "Site", "Statut", "Date", "Créé par", "Produit", "Code", "Qté reçue", "Montant (F CFA)"]
    for col, titre in enumerate(entetes, 1):
        cell = ws.cell(row=1, column=col, value=titre)
        _style_entete(cell)
    ws.row_dimensions[1].height = 30

    row = 2
    for rec in qs:
        fourn = rec.commande.fournisseur.raison_sociale if rec.commande_id else (rec.fournisseur.raison_sociale if rec.fournisseur_id else "—")
        site = rec.commande.site_destination.nom if rec.commande_id else (rec.site_destination.nom if rec.site_destination_id else "—")
        commande_ref = f"CF #{rec.commande_id}" if rec.commande_id else "Directe"
        base = [
            rec.pk,
            commande_ref,
            fourn,
            site,
            rec.get_statut_display(),
            rec.cree_le.strftime("%d/%m/%Y"),
            rec.cree_par.get_full_name() or rec.cree_par.username if rec.cree_par_id else "—",
        ]
        lignes = list(rec.lignes.all())
        if not lignes:
            for col, val in enumerate(base, 1):
                ws.cell(row=row, column=col, value=val)
            row += 1
        else:
            for ligne in lignes:
                for col, val in enumerate(base, 1):
                    ws.cell(row=row, column=col, value=val)
                ws.cell(row=row, column=8, value=ligne.produit.designation)
                ws.cell(row=row, column=9, value=ligne.produit.code)
                ws.cell(row=row, column=10, value=ligne.quantite_recue)
                ws.cell(row=row, column=11, value=float(ligne.montant))
                row += 1

    _ajuster_colonnes(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nom = f"receptions_fournisseur_{timezone.localdate().isoformat()}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nom}"'
    return response
