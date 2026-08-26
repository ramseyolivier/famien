from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import F, IntegerField, OuterRef, Q, Subquery, Sum
from django.db.models.functions import Coalesce
from django.http import HttpResponseForbidden
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone

from catalogue.models import Produit
from core.models import Profil, Site, TypeSite

from core.models import Commune
from .models import Ajustement, AjustementLigne, MouvementStock, SoldeStock, StatutAjustement
from .services import valider_ajustement, references_sous_seuil, dates_passage_sous_seuil, references_en_rupture, dates_passage_rupture


def _nb_nc_zone(u, NonConformite, StatutNonConformite):
    from django.db.models import Q
    qs = NonConformite.objects.filter(
        statut__in=[StatutNonConformite.OUVERTE, StatutNonConformite.EN_COURS]
    )
    if u.acces_national or u.is_superuser:
        return qs.count()
    q = Q(cree_par=u)
    if u.profil == Profil.SUPERVISEUR and u.commune_id:
        q |= Q(
            cree_par__profil__in=[Profil.COMMERCIAL, Profil.CHEF_EQUIPE, Profil.GEST_MAGASIN],
            cree_par__site__commune_id=u.commune_id,
        )
    elif u.profil == Profil.GEST_MAGASIN and u.site_id:
        commune_id = u.site.commune_id
        if commune_id:
            q |= Q(
                cree_par__profil__in=[Profil.COMMERCIAL, Profil.CHEF_EQUIPE],
                cree_par__site__commune_id=commune_id,
            )
    elif u.profil == Profil.CHEF_EQUIPE and u.site_id:
        q |= Q(
            cree_par__profil=Profil.COMMERCIAL,
            cree_par__site_id=u.site_id,
        )
    return qs.filter(q).count()


@login_required
def hub_stock(request):
    from ventes.views import _sites_perimetre
    from transferts.models import Transfert, StatutTransfert
    from inventaires.models import Inventaire, StatutInventaire
    from nonconformites.models import NonConformite, StatutNonConformite
    from achats.models import CommandeFournisseur, StatutCommande
    from approvisionnement.models import (
        CommandeEcole, CommandeMagasin,
        StatutCommandeEcole, StatutCommandeMagasin,
        StockReserve, StockReserveDepot,
    )
    from django.db.models import Sum as _Sum

    u = request.user
    sites = _sites_perimetre(u)

    # Kits non constructibles : au moins une ligne avec stock < quantité requise
    if u.profil != Profil.GEST_MAGASIN:
        from ventes.views import _ecoles_perimetre
        from kits.models import Kit, KitLigne
        ecoles = _ecoles_perimetre(u).filter(actif=True)
        stock_sq = SoldeStock.objects.filter(
            site=OuterRef(OuterRef("ecole")),
            produit=OuterRef("produit"),
        ).values("quantite")[:1]
        lignes_insuffisantes = KitLigne.objects.annotate(
            qt_stock=Coalesce(Subquery(stock_sq, output_field=IntegerField()), 0)
        ).filter(kit__ecole__in=ecoles, kit__actif=True, quantite__gt=F("qt_stock"))
        nb_kits_nc = Kit.objects.filter(pk__in=lignes_insuffisantes.values("kit_id"), actif=True).count()
    else:
        nb_kits_nc = 0

    return render(request, "stock/hub.html", {
        "nb_ruptures": SoldeStock.objects.filter(site__in=sites, quantite=0).count(),
        "nb_alertes": references_sous_seuil(sites=sites).count(),
        "nb_transferts": Transfert.objects.filter(
            Q(site_origine__in=sites) | Q(site_destination__in=sites),
            statut=StatutTransfert.EN_ATTENTE,
        ).count(),
        "nb_inventaires": Inventaire.objects.filter(site__in=sites, statut=StatutInventaire.EN_COURS).count(),
        "nb_nc": _nb_nc_zone(u, NonConformite, StatutNonConformite),
        "nb_kits_nc": nb_kits_nc,
        "nb_commandes": CommandeFournisseur.objects.filter(
            site_destination__in=sites,
            statut__in=[StatutCommande.SOUMIS, StatutCommande.VALIDE_N1],
        ).count(),
        # Approvisionnement école
        "nb_commandes_ecole_a_valider": CommandeEcole.objects.filter(
            statut=StatutCommandeEcole.SOUMISE
        ).count(),
        "nb_livraisons_ecole_a_preparer": (
            CommandeEcole.objects.filter(
                statut=StatutCommandeEcole.VALIDEE,
                ecole__magasin_rattachement=u.site,
            ).count()
            if getattr(u, "site_id", None)
            else 0
        ),
        "nb_receptions_ecole_a_faire": (
            CommandeEcole.objects.filter(
                statut=StatutCommandeEcole.LIVREE,
                ecole=u.site,
            ).count()
            if getattr(u, "site_id", None)
            else 0
        ),
        "nb_produits_reserves_magasin": StockReserve.objects.values("produit_id").distinct().count(),
        # Approvisionnement magasin
        "nb_commandes_magasin_a_valider": CommandeMagasin.objects.filter(
            statut=StatutCommandeMagasin.SOUMISE
        ).count(),
        "nb_livraisons_magasin_a_preparer": CommandeMagasin.objects.filter(
            statut=StatutCommandeMagasin.VALIDEE
        ).count(),
        "nb_livraisons_magasin_a_recevoir": (
            CommandeMagasin.objects.filter(
                statut=StatutCommandeMagasin.LIVREE,
                magasin=u.site,
            ).count()
            if getattr(u, "site_id", None)
            else 0
        ),
        "nb_produits_reserves_depot": StockReserveDepot.objects.values("produit_id").distinct().count(),
    })


# ─── Références sous seuil de sécurité ──────────────────────────────────────

@login_required
def stock_sous_seuil(request):
    from ventes.views import _sites_perimetre
    u = request.user
    sites = _sites_perimetre(u)

    filtre_commune  = request.GET.get("commune", "")
    filtre_type     = request.GET.get("type_site", "")
    filtre_produit  = request.GET.get("produit", "")

    qs = references_sous_seuil(sites=sites).select_related("site__commune", "produit")

    if filtre_commune:
        qs = qs.filter(site__commune_id=filtre_commune)
    if filtre_type:
        qs = qs.filter(site__type=filtre_type)
    if filtre_produit:
        qs = qs.filter(produit_id=filtre_produit)

    site_ids = list(sites.values_list("pk", flat=True))
    dates = dates_passage_sous_seuil(site_ids)

    from django.utils import timezone
    aujourd_hui = timezone.now().date()

    lignes = []
    for s in qs:
        date_p = dates.get((s.site_id, s.produit_id))
        jours = (aujourd_hui - date_p.date()).days if date_p else None
        lignes.append({
            "solde": s,
            "jours": jours,
            "en_rupture": s.quantite == 0,
        })

    communes = (
        Commune.objects.filter(sites__in=sites).distinct().order_by("nom")
        if u.acces_national or u.is_superuser else None
    )
    produits_filtre = (
        SoldeStock.objects.filter(site__in=sites, quantite__lte=F("stock_securite"))
        .select_related("produit").order_by("produit__code")
        .values("produit_id", "produit__code", "produit__designation").distinct()
    )

    return render(request, "stock/sous_seuil.html", {
        "lignes": lignes,
        "communes": communes,
        "types_site": TypeSite.choices,
        "produits_filtre": produits_filtre,
        "filtre_commune": filtre_commune,
        "filtre_type": filtre_type,
        "filtre_produit": filtre_produit,
        "nb_total": len(lignes),
    })


@login_required
def stock_ruptures(request):
    from ventes.views import _sites_perimetre
    u = request.user
    sites = _sites_perimetre(u)

    filtre_commune = request.GET.get("commune", "")
    filtre_type    = request.GET.get("type_site", "")
    filtre_produit = request.GET.get("produit", "")

    qs = references_en_rupture(sites=sites)
    if filtre_commune:
        qs = qs.filter(site__commune_id=filtre_commune)
    if filtre_type:
        qs = qs.filter(site__type=filtre_type)
    if filtre_produit:
        qs = qs.filter(produit_id=filtre_produit)

    site_ids = list(sites.values_list("pk", flat=True))
    dates_map = dates_passage_rupture(site_ids)

    lignes = []
    for s in qs.select_related("site", "site__commune", "produit"):
        depuis = dates_map.get((s.site_id, s.produit_id))
        jours = (timezone.now() - depuis).days if depuis else None
        lignes.append({"solde": s, "jours": jours})

    communes = (
        Commune.objects.filter(sites__in=sites).distinct().order_by("nom")
        if u.acces_national or u.is_superuser else None
    )
    produits_filtre = (
        SoldeStock.objects.filter(site__in=sites, quantite=0)
        .select_related("produit").order_by("produit__code")
        .values("produit_id", "produit__code", "produit__designation").distinct()
    )

    return render(request, "stock/ruptures.html", {
        "lignes": lignes,
        "communes": communes,
        "types_site": TypeSite.choices,
        "produits_filtre": produits_filtre,
        "filtre_commune": filtre_commune,
        "filtre_type": filtre_type,
        "filtre_produit": filtre_produit,
        "nb_total": len(lignes),
    })


# ─── Ajustements de stock (M12) ──────────────────────────────────────────────

@login_required
def ajustements_liste(request):
    sites = request.user.sites_autorises()
    qs = (
        Ajustement.objects
        .filter(site__in=sites)
        .select_related("site", "cree_par")
        .order_by("-cree_le")[:100]
    )
    return render(request, "stock/ajustements_liste.html", {"ajustements": qs})


@login_required
def ajustement_formulaire(request):
    sites = request.user.sites_autorises()
    produits = Produit.objects.filter(actif=True).order_by("code")

    if request.method == "POST":
        site_id = request.POST.get("site")
        motif = request.POST.get("motif", "").strip()
        produit_ids = request.POST.getlist("produit_id")
        quantites = request.POST.getlist("quantite")

        if not sites.filter(pk=site_id).exists():
            messages.error(request, "Site non autorisé.")
        elif not motif:
            messages.error(request, "Le motif est obligatoire.")
        else:
            lignes = []
            for pid, q in zip(produit_ids, quantites):
                try:
                    q_int = int(q)
                    if q_int != 0:
                        lignes.append((int(pid), q_int))
                except (ValueError, TypeError):
                    pass
            if not lignes:
                messages.error(request, "Ajoutez au moins une ligne avec une quantité non nulle.")
            else:
                with transaction.atomic():
                    aju = Ajustement.objects.create(site_id=site_id, motif=motif, cree_par=request.user)
                    AjustementLigne.objects.bulk_create([
                        AjustementLigne(ajustement=aju, produit_id=pid, quantite=q)
                        for pid, q in lignes
                    ])
                messages.success(request, "Ajustement créé en brouillon.")
                return redirect("ajustements_liste")

    sites_list = list(sites)
    site_unique = sites_list[0] if len(sites_list) == 1 else None
    return render(request, "stock/ajustement_formulaire.html", {
        "sites": sites_list,
        "site_unique": site_unique,
        "produits": produits,
    })


@login_required
def ajustement_valider(request, pk):
    sites = request.user.sites_autorises()
    aju = get_object_or_404(Ajustement, pk=pk, statut=StatutAjustement.BROUILLON)
    if not sites.filter(pk=aju.site_id).exists():
        messages.error(request, "Accès refusé.")
        return redirect("ajustements_liste")
    if request.method == "POST":
        try:
            valider_ajustement(aju, par=request.user)
            messages.success(request, "Ajustement validé — stock mis à jour.")
        except Exception as e:
            messages.error(request, str(e))
    return redirect("ajustements_liste")


# ─── Paramétrage des seuils (M07.2) — DG et MANAGER uniquement ───────────────

@login_required
def parametres_seuils(request):
    u = request.user
    if u.profil not in (Profil.DG, Profil.MANAGER) and not u.is_superuser:
        return HttpResponseForbidden("Accès réservé au DG et au Manager.")

    site_id = request.GET.get("site") or request.POST.get("site_filtre")
    sites_qs = Site.objects.filter(
        type__in=[TypeSite.ECOLE, TypeSite.MAGASIN, TypeSite.DEPOT]
    ).order_by("type", "nom")

    site = None
    rows = []  # liste de (produit, solde_ou_None)

    if site_id:
        site = get_object_or_404(Site, pk=site_id)
        produits = Produit.objects.filter(actif=True).order_by("designation")
        soldes_map = {
            s.produit_id: s
            for s in SoldeStock.objects.filter(site=site)
        }
        rows = [(p, soldes_map.get(p.pk)) for p in produits]

    if request.method == "POST" and "enregistrer" in request.POST and site:
        erreurs = []
        parsed = {}
        for produit, solde in rows:
            mini_s = request.POST.get(f"mini_{produit.pk}", "").strip()
            seuil_s = request.POST.get(f"seuil_{produit.pk}", "").strip()
            maxi_s = request.POST.get(f"maxi_{produit.pk}", "").strip()
            try:
                new_mini = max(0, int(mini_s)) if mini_s else 0
                new_seuil = max(0, int(seuil_s)) if seuil_s else 0
                new_maxi = max(0, int(maxi_s)) if maxi_s else 0
            except ValueError:
                continue
            if new_seuil < new_mini:
                erreurs.append(f"{produit.code} : seuil d'alerte ({new_seuil}) < mini ({new_mini})")
            elif new_seuil > 0 and new_maxi <= new_seuil:
                erreurs.append(f"{produit.code} : max ({new_maxi}) doit être strictement supérieur au seuil d'alerte ({new_seuil})")
            else:
                parsed[produit.pk] = (new_mini, new_seuil, new_maxi, solde)

        if erreurs:
            for e in erreurs:
                messages.error(request, e)
        else:
            updated = 0
            with transaction.atomic():
                for produit, solde in rows:
                    if produit.pk not in parsed:
                        continue
                    new_mini, new_seuil, new_maxi, solde = parsed[produit.pk]
                    if solde:
                        if (new_mini, new_seuil, new_maxi) != (solde.stock_minimum, solde.stock_securite, solde.stock_maximum):
                            SoldeStock.objects.filter(pk=solde.pk).update(
                                stock_minimum=new_mini,
                                stock_securite=new_seuil,
                                stock_maximum=new_maxi,
                            )
                            updated += 1
                    elif new_mini or new_seuil or new_maxi:
                        SoldeStock.objects.create(
                            site=site, produit=produit, quantite=0,
                            stock_minimum=new_mini,
                            stock_securite=new_seuil,
                            stock_maximum=new_maxi,
                        )
                        updated += 1
            if updated:
                messages.success(request, f"{updated} ligne{'s' if updated > 1 else ''} mise{'s' if updated > 1 else ''} à jour.")
            else:
                messages.info(request, "Aucune modification détectée.")
            return redirect(f"/stock/seuils/?site={site_id}")

    return render(request, "stock/seuils.html", {
        "rows": rows,
        "sites": sites_qs,
        "site": site,
        "site_filtre": site_id,
    })


def _construire_rapport(user, site_id, commune_id, debut, fin, produit_id=""):
    """Retourne (lignes, meta) pour la période [debut, fin] sur les sites autorisés."""
    from collections import defaultdict
    from approvisionnement.models import StockReserve, StockReserveDepot

    est_dg_manager  = user.is_superuser or user.profil in {Profil.DG, Profil.MANAGER}
    est_superviseur = user.profil == Profil.SUPERVISEUR

    sites_base = user.sites_autorises()
    sites_qs   = sites_base
    if site_id:
        sites_qs = sites_qs.filter(pk=site_id)
    if commune_id and est_dg_manager:
        sites_qs = sites_qs.filter(commune_id=commune_id)

    site_ids   = list(sites_qs.values_list("pk", flat=True))
    sites_list = list(sites_qs.select_related("commune"))

    # Mouvements de la période groupés par (site, produit, type)
    mvts_filtre = dict(site_id__in=site_ids, horodatage__date__gte=debut, horodatage__date__lte=fin)
    if produit_id:
        mvts_filtre["produit_id"] = produit_id
    mvts_periode_qs = (
        MouvementStock.objects
        .filter(**mvts_filtre)
        .values("site_id", "produit_id", "type")
        .annotate(total=Sum("quantite"))
    )

    # Mouvements APRÈS la fin de période (pour reconstituer la clôture)
    apres_index = defaultdict(int)
    for r in (
        MouvementStock.objects
        .filter(site_id__in=site_ids, horodatage__date__gt=fin)
        .values("site_id", "produit_id")
        .annotate(t=Sum("quantite"))
    ):
        apres_index[(r["site_id"], r["produit_id"])] = r["t"]

    mvt_par_type    = defaultdict(lambda: defaultdict(int))
    mvt_total_periode = defaultdict(int)
    for m in mvts_periode_qs:
        key = (m["site_id"], m["produit_id"])
        mvt_par_type[key][m["type"]] = m["total"]
        mvt_total_periode[key] += m["total"]

    soldes_qs = SoldeStock.objects.filter(site_id__in=site_ids)
    if produit_id:
        soldes_qs = soldes_qs.filter(produit_id=produit_id)
    soldes = soldes_qs.select_related("produit", "site", "site__commune")

    site_types      = {s.type for s in sites_list}
    afficher_vendue  = TypeSite.ECOLE in site_types or est_dg_manager
    afficher_livree  = TypeSite.MAGASIN in site_types or TypeSite.DEPOT in site_types or est_dg_manager
    afficher_reservee = TypeSite.MAGASIN in site_types or TypeSite.DEPOT in site_types or est_dg_manager
    afficher_site    = sites_base.count() > 1  # basé sur le périmètre réel, pas le filtre

    # Réservations actuelles par (magasin, produit)
    magasin_ids = [s.pk for s in sites_list if s.type == TypeSite.MAGASIN]
    depot_ids   = [s.pk for s in sites_list if s.type == TypeSite.DEPOT]
    reserves_index = defaultdict(int)
    if magasin_ids:
        for r in (
            StockReserve.objects
            .filter(magasin_id__in=magasin_ids)
            .values("magasin_id", "produit_id")
            .annotate(total=Sum("quantite"))
        ):
            reserves_index[(r["magasin_id"], r["produit_id"])] = r["total"]
    if depot_ids:
        for r in StockReserveDepot.objects.values("produit_id").annotate(total=Sum("quantite")):
            for depot_id in depot_ids:
                reserves_index[(depot_id, r["produit_id"])] = r["total"]

    TYPES_RECUS = {"ENTREE_ACHAT", "ENTREE_LIVRAISON_ECOLE", "RETOUR_LIVRAISON_ECOLE", "ENTREE_APPRO_MAGASIN"}

    lignes = []
    for s in soldes:
        key   = (s.site_id, s.produit_id)
        types = mvt_par_type[key]

        cloture        = s.quantite - apres_index[key]
        total_periode  = mvt_total_periode[key]
        initiale       = cloture - total_periode

        recue     = sum(v for k, v in types.items() if k in TYPES_RECUS)
        transfert = types.get("ENTREE_TRANSFERT", 0) + types.get("SORTIE_TRANSFERT", 0)
        vendue    = max(0, abs(types.get("SORTIE_VENTE", 0)) - types.get("ANNULATION", 0))
        livree    = abs(types.get("SORTIE_LIVRAISON_ECOLE", 0)) + abs(types.get("SORTIE_APPRO_MAGASIN", 0))
        ajuste   = types.get("AJUSTEMENT", 0)
        reservee = reserves_index.get(key, 0)

        if s.quantite == 0 and not types:
            continue

        lignes.append({
            "site":      s.site,
            "produit":   s.produit,
            "initiale":  initiale,
            "recue":     recue,
            "transfert": transfert,
            "vendue":    vendue,
            "livree":    livree,
            "ajuste":    ajuste,
            "cloture":   cloture,
            "reservee":  reservee,
        })

    meta = {
        "afficher_vendue":   afficher_vendue,
        "afficher_livree":   afficher_livree,
        "afficher_reservee": afficher_reservee,
        "afficher_site":     afficher_site,
        "sites_base":        sites_base,
        "est_dg_manager":    est_dg_manager,
        "est_superviseur":   est_superviseur,
    }
    return lignes, meta


@login_required
def rapport_journalier(request):
    from datetime import date as date_type

    user  = request.user
    today = timezone.localdate()

    debut_str  = request.GET.get("debut", "")
    fin_str    = request.GET.get("fin", "")
    site_id    = request.GET.get("site", "")
    commune_id = request.GET.get("commune", "")
    produit_id = request.GET.get("produit", "")
    tri        = request.GET.get("tri", "produit")
    sens       = request.GET.get("sens", "asc")

    try:
        debut = date_type.fromisoformat(debut_str) if debut_str else today
    except ValueError:
        debut = today
    try:
        fin = date_type.fromisoformat(fin_str) if fin_str else today
    except ValueError:
        fin = today
    if fin < debut:
        fin = debut

    est_dg_manager  = user.is_superuser or user.profil in {Profil.DG, Profil.MANAGER}
    est_superviseur = user.profil == Profil.SUPERVISEUR

    lignes, meta = _construire_rapport(user, site_id, commune_id, debut, fin, produit_id)

    CHAMPS = {
        "produit":     lambda r: r["produit"].designation,
        "site":        lambda r: r["site"].nom,
        "initiale":    lambda r: r["initiale"],
        "recue":       lambda r: r["recue"],
        "transfert":   lambda r: r["transfert"],
        "vendue":  lambda r: r["vendue"],
        "livree":  lambda r: r["livree"],
        "ajuste":   lambda r: r["ajuste"],
        "cloture":  lambda r: r["cloture"],
        "reservee": lambda r: r["reservee"],
    }
    if tri in CHAMPS:
        lignes.sort(key=CHAMPS[tri], reverse=(sens == "desc"))

    communes_filtre = Commune.objects.order_by("nom") if est_dg_manager else None
    if est_dg_manager or est_superviseur:
        sites_filtre_qs = meta["sites_base"]
        if commune_id:
            sites_filtre_qs = sites_filtre_qs.filter(commune_id=commune_id)
        sites_filtre = sites_filtre_qs.order_by("type", "nom")
    else:
        sites_filtre = None

    produits_filtre = Produit.objects.filter(actif=True).order_by("designation")

    filtres = {
        "debut":    str(debut),
        "fin":      str(fin),
        "site":     site_id,
        "commune":  commune_id,
        "produit":  produit_id,
        "tri":      tri,
        "sens":     sens,
    }

    return render(request, "stock/rapport_journalier.html", {
        "lignes":           lignes,
        "debut":            debut,
        "fin":              fin,
        "filtres":          filtres,
        "communes_filtre":  communes_filtre,
        "sites_filtre":     sites_filtre,
        "produits_filtre":  produits_filtre,
        **meta,
    })


@login_required
def rapport_journalier_export(request):
    import io
    import openpyxl
    from datetime import date as date_type
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from django.http import HttpResponse

    user  = request.user
    today = timezone.localdate()

    debut_str  = request.GET.get("debut", "")
    fin_str    = request.GET.get("fin", "")
    site_id    = request.GET.get("site", "")
    commune_id = request.GET.get("commune", "")
    produit_id = request.GET.get("produit", "")

    try:
        debut = date_type.fromisoformat(debut_str) if debut_str else today
    except ValueError:
        debut = today
    try:
        fin = date_type.fromisoformat(fin_str) if fin_str else today
    except ValueError:
        fin = today
    if fin < debut:
        fin = debut

    lignes, meta = _construire_rapport(user, site_id, commune_id, debut, fin, produit_id)
    lignes.sort(key=lambda r: (r["site"].nom, r["produit"].designation))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rapport journalier"

    entetes = ["Site", "Code", "Désignation", "Initiale", "Reçue", "Transféré"]
    if meta["afficher_vendue"]:
        entetes.append("Vendue")
    if meta["afficher_livree"]:
        entetes.append("Livrée")
    entetes += ["Ajustée", "Final", "Réservée"]

    for col, titre in enumerate(entetes, 1):
        c = ws.cell(row=1, column=col, value=titre)
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor="16233F")
        c.alignment = Alignment(horizontal="center", vertical="center")

    for row, l in enumerate(lignes, 2):
        col = 1
        ws.cell(row=row, column=col, value=l["site"].nom);      col += 1
        ws.cell(row=row, column=col, value=l["produit"].code);  col += 1
        ws.cell(row=row, column=col, value=l["produit"].designation); col += 1
        ws.cell(row=row, column=col, value=l["initiale"]);      col += 1
        ws.cell(row=row, column=col, value=l["recue"] or None); col += 1
        ws.cell(row=row, column=col, value=l["transfert"] or None); col += 1
        if meta["afficher_vendue"]:
            ws.cell(row=row, column=col, value=-l["vendue"] if l["vendue"] else None); col += 1
        if meta["afficher_livree"]:
            ws.cell(row=row, column=col, value=-l["livree"] if l["livree"] else None); col += 1
        ws.cell(row=row, column=col, value=l["ajuste"] or None);   col += 1
        ws.cell(row=row, column=col, value=l["cloture"]);           col += 1
        ws.cell(row=row, column=col, value=l["reservee"] or None)

    for col in ws.columns:
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(
            max(len(str(c.value or "")) for c in col) + 4, 50
        )

    nom = f"rapport_{debut}_{fin}.xlsx"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(buf, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nom}"'
    return response
