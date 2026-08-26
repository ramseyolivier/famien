"""Exports Excel pour l'état du stock (M24)."""

import io

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.utils import timezone

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from stock.models import SoldeStock
from stock.services import references_sous_seuil, dates_passage_sous_seuil, references_en_rupture, dates_passage_rupture


def _style_entete(cell):
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.fill = PatternFill("solid", fgColor="16233F")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _ajuster_colonnes(ws):
    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 50)


@login_required
def export_stock_excel(request):
    from ventes.views import _sites_perimetre

    sites = _sites_perimetre(request.user)
    soldes = (
        SoldeStock.objects
        .filter(site__in=sites)
        .select_related("site", "produit", "produit__categorie")
        .order_by("produit__code", "site__nom")
    )

    site_id = request.GET.get("site", "")
    if site_id.isdigit():
        soldes = soldes.filter(site_id=site_id)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stock"
    ws.freeze_panes = "A2"

    entetes = ["Code", "Désignation", "Catégorie", "Site", "Quantité", "Seuil alerte", "Statut"]
    for col, titre in enumerate(entetes, 1):
        cell = ws.cell(row=1, column=col, value=titre)
        _style_entete(cell)
    ws.row_dimensions[1].height = 30

    for row, s in enumerate(soldes, 2):
        ws.cell(row=row, column=1, value=s.produit.code)
        ws.cell(row=row, column=2, value=s.produit.designation)
        ws.cell(row=row, column=3, value=s.produit.categorie.nom if s.produit.categorie_id else "")
        ws.cell(row=row, column=4, value=s.site.nom)
        ws.cell(row=row, column=5, value=s.quantite)
        ws.cell(row=row, column=6, value=s.stock_securite)
        statut = "Rupture" if s.quantite == 0 else ("Alerte" if s.quantite <= s.stock_securite else "OK")
        ws.cell(row=row, column=7, value=statut)

    _ajuster_colonnes(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nom = f"stock_{timezone.localdate().isoformat()}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nom}"'
    return response


@login_required
def export_sous_seuil_excel(request):
    from ventes.views import _sites_perimetre
    from datetime import date

    u = request.user
    sites = _sites_perimetre(u)
    soldes = references_sous_seuil(sites=sites).select_related("site", "site__commune", "produit")
    site_ids = list(sites.values_list("pk", flat=True))
    dates_map = dates_passage_sous_seuil(site_ids)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sous le seuil"
    ws.freeze_panes = "A2"

    entetes = ["Commune", "Type de site", "Site", "Code produit", "Désignation", "En stock", "Seuil", "Jours sous seuil"]
    for col, titre in enumerate(entetes, 1):
        cell = ws.cell(row=1, column=col, value=titre)
        _style_entete(cell)
    ws.row_dimensions[1].height = 30

    rouge = PatternFill("solid", fgColor="FEE2E2")
    orange = PatternFill("solid", fgColor="FEF3C7")

    for row, s in enumerate(soldes, 2):
        depuis = dates_map.get((s.site_id, s.produit_id))
        jours = (timezone.now() - depuis).days if depuis else None

        ws.cell(row=row, column=1, value=s.site.commune.nom if s.site.commune_id else "—")
        ws.cell(row=row, column=2, value=s.site.get_type_display())
        ws.cell(row=row, column=3, value=s.site.nom)
        ws.cell(row=row, column=4, value=s.produit.code)
        ws.cell(row=row, column=5, value=s.produit.designation)
        ws.cell(row=row, column=6, value=int(s.quantite))
        ws.cell(row=row, column=7, value=int(s.stock_securite))
        cell_j = ws.cell(row=row, column=8, value=jours)
        if jours is not None:
            if jours >= 7:
                cell_j.fill = rouge
            elif jours >= 3:
                cell_j.fill = orange

    _ajuster_colonnes(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nom = f"sous_seuil_{timezone.localdate().isoformat()}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nom}"'
    return response


@login_required
def export_ruptures_excel(request):
    from ventes.views import _sites_perimetre

    u = request.user
    sites = _sites_perimetre(u)
    soldes = references_en_rupture(sites=sites).select_related("site", "site__commune", "produit")
    site_ids = list(sites.values_list("pk", flat=True))
    dates_map = dates_passage_rupture(site_ids)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ruptures"
    ws.freeze_panes = "A2"

    entetes = ["Commune", "Type de site", "Site", "Code produit", "Désignation", "En stock", "Seuil", "Depuis quand (jours)"]
    for col, titre in enumerate(entetes, 1):
        cell = ws.cell(row=1, column=col, value=titre)
        _style_entete(cell)
    ws.row_dimensions[1].height = 30

    rouge = PatternFill("solid", fgColor="FEE2E2")
    orange = PatternFill("solid", fgColor="FEF3C7")

    for row, s in enumerate(soldes, 2):
        depuis = dates_map.get((s.site_id, s.produit_id))
        jours = (timezone.now() - depuis).days if depuis else None

        ws.cell(row=row, column=1, value=s.site.commune.nom if s.site.commune_id else "—")
        ws.cell(row=row, column=2, value=s.site.get_type_display())
        ws.cell(row=row, column=3, value=s.site.nom)
        ws.cell(row=row, column=4, value=s.produit.code)
        ws.cell(row=row, column=5, value=s.produit.designation)
        ws.cell(row=row, column=6, value=0)
        ws.cell(row=row, column=7, value=int(s.stock_securite))
        cell_j = ws.cell(row=row, column=8, value=jours)
        if jours is not None:
            if jours >= 7:
                cell_j.fill = rouge
            elif jours >= 3:
                cell_j.fill = orange

    _ajuster_colonnes(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nom = f"ruptures_{timezone.localdate().isoformat()}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nom}"'
    return response
