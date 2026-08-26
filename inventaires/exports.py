"""Export Excel d'un inventaire validé (M24)."""

import io

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from inventaires.models import Inventaire


def _style_entete(cell):
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.fill = PatternFill("solid", fgColor="16233F")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _ajuster_colonnes(ws):
    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 50)


@login_required
def export_inventaire_excel(request, pk):
    inv = get_object_or_404(Inventaire, pk=pk)
    if not request.user.peut_acceder_au_site(inv.site):
        from django.http import HttpResponseForbidden
        return HttpResponseForbidden()

    lignes = inv.lignes.select_related("produit", "produit__categorie").order_by("produit__code")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Inventaire {inv.pk}"
    ws.freeze_panes = "A2"

    entetes = ["Code", "Désignation", "Catégorie", "Théorique", "Physique", "Écart", "Commentaire"]
    for col, titre in enumerate(entetes, 1):
        cell = ws.cell(row=1, column=col, value=titre)
        _style_entete(cell)

    for row, l in enumerate(lignes, 2):
        ws.cell(row=row, column=1, value=l.produit.code)
        ws.cell(row=row, column=2, value=l.produit.designation)
        ws.cell(row=row, column=3, value=l.produit.categorie.nom if l.produit.categorie_id else "")
        ws.cell(row=row, column=4, value=l.quantite_theorique)
        ws.cell(row=row, column=5, value=l.quantite_physique)
        ws.cell(row=row, column=6, value=l.ecart)
        ws.cell(row=row, column=7, value=l.commentaire or "")

    _ajuster_colonnes(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nom = f"inventaire_{inv.pk}_{timezone.localdate().isoformat()}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nom}"'
    return response
