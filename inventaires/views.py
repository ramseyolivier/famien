from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.shortcuts import get_object_or_404, redirect, render

from core.models import Commune, Profil, Site, TypeSite
from stock.models import SoldeStock

from .models import Inventaire, InventaireLigne, StatutInventaire
from .services import valider_inventaire, verifier_blocages_inventaire

PEUT_INITIER = {Profil.DG, Profil.MANAGER, Profil.CHEF_EQUIPE, Profil.GEST_MAGASIN}
_DG_MANAGER = {Profil.DG, Profil.MANAGER}


def _peut_initier(user):
    return user.is_superuser or user.profil in PEUT_INITIER


def _est_dg_manager(user):
    return user.is_superuser or user.profil in _DG_MANAGER


@login_required
def inventaires_liste(request):
    u = request.user
    est_dg = _est_dg_manager(u)

    if est_dg:
        # Le DG/Manager voit uniquement les dépôts (son site_id s'il est défini, sinon tous les DEPOT)
        if u.site_id:
            sites = Site.objects.filter(pk=u.site_id)
        else:
            sites = Site.objects.filter(type=TypeSite.DEPOT)
    else:
        sites = u.sites_autorises()

    est_superviseur = u.profil == Profil.SUPERVISEUR
    statut_filtre = request.GET.get("statut", "")
    site_filtre   = request.GET.get("site", "")
    debut         = request.GET.get("debut", "")
    fin           = request.GET.get("fin", "")

    qs = (
        Inventaire.objects
        .filter(site__in=sites)
        .select_related("site", "cree_par")
        .order_by("-cree_le")
    )
    if site_filtre and est_superviseur:
        qs = qs.filter(site_id=site_filtre)
    if statut_filtre:
        qs = qs.filter(statut=statut_filtre)
    if debut:
        qs = qs.filter(cree_le__date__gte=debut)
    if fin:
        qs = qs.filter(cree_le__date__lte=fin)

    sites_filtre = sites.order_by("nom") if est_superviseur else None
    filtres = {"statut": statut_filtre, "site": site_filtre, "debut": debut, "fin": fin}

    return render(request, "inventaires/liste.html", {
        "inventaires":    qs[:200],
        "peut_initier":   _peut_initier(u),
        "est_dg_manager": est_dg,
        "est_superviseur": est_superviseur,
        "statuts":        StatutInventaire.choices,
        "filtres":        filtres,
        "sites_filtre":   sites_filtre,
        "depot":          u.site if est_dg else None,
    })


@login_required
def inventaires_liste_export(request):
    """Export Excel de la liste filtrée de /inventaires/."""
    import io
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from django.http import HttpResponse
    from django.utils import timezone

    u = request.user
    est_dg = _est_dg_manager(u)
    if est_dg:
        if u.site_id:
            sites = Site.objects.filter(pk=u.site_id)
        else:
            sites = Site.objects.filter(type=TypeSite.DEPOT)
    else:
        sites = u.sites_autorises()

    statut_filtre = request.GET.get("statut", "")
    debut         = request.GET.get("debut", "")
    fin           = request.GET.get("fin", "")

    qs = Inventaire.objects.filter(site__in=sites).select_related("site", "cree_par", "valide_par").order_by("-cree_le")
    if statut_filtre:
        qs = qs.filter(statut=statut_filtre)
    if debut:
        qs = qs.filter(cree_le__date__gte=debut)
    if fin:
        qs = qs.filter(cree_le__date__lte=fin)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventaires"
    entetes = ["#", "Site", "Statut", "Écarts", "Créé le", "Par", "Validé le", "Par"]
    for col, t in enumerate(entetes, 1):
        c = ws.cell(row=1, column=col, value=t)
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor="16233F")
        c.alignment = Alignment(horizontal="center")
    for row, inv in enumerate(qs, 2):
        ws.cell(row=row, column=1, value=f"INV-{inv.pk}")
        ws.cell(row=row, column=2, value=inv.site.nom)
        ws.cell(row=row, column=3, value=inv.get_statut_display())
        ws.cell(row=row, column=4, value=inv.nb_ecarts)
        ws.cell(row=row, column=5, value=inv.cree_le.strftime("%d/%m/%Y %H:%M") if inv.cree_le else "")
        ws.cell(row=row, column=6, value=inv.cree_par.get_full_name() or inv.cree_par.username)
        ws.cell(row=row, column=7, value=inv.valide_le.strftime("%d/%m/%Y %H:%M") if inv.valide_le else "")
        ws.cell(row=row, column=8, value=inv.valide_par.get_full_name() if inv.valide_par else "")
    for col in ws.columns:
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(len(str(c.value or "")) for c in col) + 4, 40)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nom = f"inventaires_{timezone.localdate().isoformat()}.xlsx"
    resp = HttpResponse(buf, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{nom}"'
    return resp


@login_required
def inventaires_global(request):
    """Vue de supervision DG/MANAGER : tous les inventaires du système avec filtres."""
    if not _est_dg_manager(request.user):
        messages.error(request, "Accès réservé au DG et au Manager.")
        return redirect("inventaires_liste")

    statut_filtre  = request.GET.get("statut", "")
    commune_filtre = request.GET.get("commune", "")
    site_filtre    = request.GET.get("site", "")
    debut          = request.GET.get("debut", "")
    fin            = request.GET.get("fin", "")

    qs = (
        Inventaire.objects
        .select_related("site", "site__commune", "cree_par", "valide_par")
        .order_by("-cree_le")
    )
    if statut_filtre:
        qs = qs.filter(statut=statut_filtre)
    if commune_filtre:
        qs = qs.filter(site__commune_id=commune_filtre)
    if site_filtre:
        qs = qs.filter(site_id=site_filtre)
    if debut:
        qs = qs.filter(cree_le__date__gte=debut)
    if fin:
        qs = qs.filter(cree_le__date__lte=fin)

    communes_filtre = Commune.objects.order_by("nom")
    sites_filtre_qs = Site.objects.order_by("type", "nom")
    if commune_filtre:
        sites_filtre_qs = sites_filtre_qs.filter(commune_id=commune_filtre)

    filtres = {
        "statut": statut_filtre, "commune": commune_filtre,
        "site": site_filtre, "debut": debut, "fin": fin,
    }
    return render(request, "inventaires/global.html", {
        "inventaires":    qs[:500],
        "statuts":        StatutInventaire.choices,
        "communes_filtre": communes_filtre,
        "sites_filtre":   sites_filtre_qs,
        "filtres":        filtres,
    })


@login_required
def inventaires_global_export(request):
    """Export Excel de la liste filtrée de /inventaires/global/."""
    import io
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from django.http import HttpResponse
    from django.utils import timezone

    if not _est_dg_manager(request.user):
        from django.http import HttpResponseForbidden
        return HttpResponseForbidden()

    statut_filtre  = request.GET.get("statut", "")
    commune_filtre = request.GET.get("commune", "")
    site_filtre    = request.GET.get("site", "")
    debut          = request.GET.get("debut", "")
    fin            = request.GET.get("fin", "")

    qs = Inventaire.objects.select_related("site", "site__commune", "cree_par", "valide_par").order_by("-cree_le")
    if statut_filtre:
        qs = qs.filter(statut=statut_filtre)
    if commune_filtre:
        qs = qs.filter(site__commune_id=commune_filtre)
    if site_filtre:
        qs = qs.filter(site_id=site_filtre)
    if debut:
        qs = qs.filter(cree_le__date__gte=debut)
    if fin:
        qs = qs.filter(cree_le__date__lte=fin)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventaires global"
    entetes = ["#", "Commune", "Site", "Type", "Statut", "Écarts", "Créé le", "Par", "Validé le", "Par"]
    for col, t in enumerate(entetes, 1):
        c = ws.cell(row=1, column=col, value=t)
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor="16233F")
        c.alignment = Alignment(horizontal="center")
    for row, inv in enumerate(qs, 2):
        ws.cell(row=row, column=1, value=f"INV-{inv.pk}")
        ws.cell(row=row, column=2, value=inv.site.commune.nom if inv.site.commune_id else "")
        ws.cell(row=row, column=3, value=inv.site.nom)
        ws.cell(row=row, column=4, value=inv.site.get_type_display())
        ws.cell(row=row, column=5, value=inv.get_statut_display())
        ws.cell(row=row, column=6, value=inv.nb_ecarts)
        ws.cell(row=row, column=7, value=inv.cree_le.strftime("%d/%m/%Y %H:%M") if inv.cree_le else "")
        ws.cell(row=row, column=8, value=inv.cree_par.get_full_name() or inv.cree_par.username)
        ws.cell(row=row, column=9, value=inv.valide_le.strftime("%d/%m/%Y %H:%M") if inv.valide_le else "")
        ws.cell(row=row, column=10, value=inv.valide_par.get_full_name() if inv.valide_par else "")
    for col in ws.columns:
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(len(str(c.value or "")) for c in col) + 4, 40)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nom = f"inventaires_global_{timezone.localdate().isoformat()}.xlsx"
    resp = HttpResponse(buf, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{nom}"'
    return resp


@login_required
def inventaire_nouveau(request):
    if not _peut_initier(request.user):
        messages.error(request, "Vous n'êtes pas autorisé à créer un inventaire.")
        return redirect("inventaires_liste")

    u = request.user
    if _est_dg_manager(u):
        sites = Site.objects.filter(pk=u.site_id) if u.site_id else Site.objects.filter(type=TypeSite.DEPOT)
    else:
        sites = u.sites_autorises()
    sites_list = list(sites.order_by("type", "nom"))
    site_unique = sites_list[0] if len(sites_list) == 1 else None

    blocages = []
    site_verifie = None

    if request.method == "POST":
        site_id = request.POST.get("site")
        if not sites.filter(pk=site_id).exists():
            messages.error(request, "Site non autorisé.")
        else:
            site_obj = sites.select_related().get(pk=site_id)
            en_cours = Inventaire.objects.filter(site=site_obj, statut=StatutInventaire.EN_COURS).first()
            if en_cours:
                messages.info(request, "Un inventaire est déjà en cours pour ce site.")
                return redirect("inventaire_saisir", pk=en_cours.pk)
            blocages = verifier_blocages_inventaire(site_obj)
            if blocages:
                site_verifie = site_obj
            else:
                inv = Inventaire.objects.create(site_id=site_id, cree_par=request.user)
                soldes = SoldeStock.objects.filter(site_id=site_id).select_related("produit")
                InventaireLigne.objects.bulk_create([
                    InventaireLigne(
                        inventaire=inv,
                        produit=s.produit,
                        quantite_theorique=s.quantite,
                        quantite_physique=s.quantite,
                    )
                    for s in soldes
                ])
                messages.success(request, "Inventaire créé. Saisissez les quantités physiques.")
                return redirect("inventaire_saisir", pk=inv.pk)

    elif site_unique:
        en_cours = Inventaire.objects.filter(site=site_unique, statut=StatutInventaire.EN_COURS).first()
        if en_cours:
            messages.info(request, "Un inventaire est déjà en cours pour ce site.")
            return redirect("inventaire_saisir", pk=en_cours.pk)
        blocages = verifier_blocages_inventaire(site_unique)
        if not blocages:
            inv = Inventaire.objects.create(site=site_unique, cree_par=request.user)
            soldes = SoldeStock.objects.filter(site=site_unique).select_related("produit")
            InventaireLigne.objects.bulk_create([
                InventaireLigne(
                    inventaire=inv,
                    produit=s.produit,
                    quantite_theorique=s.quantite,
                    quantite_physique=s.quantite,
                )
                for s in soldes
            ])
            return redirect("inventaire_saisir", pk=inv.pk)
        site_verifie = site_unique

    return render(request, "inventaires/nouveau.html", {
        "sites": sites_list,
        "site_unique": site_unique,
        "blocages": blocages,
        "site_verifie": site_verifie,
    })


@login_required
def inventaire_saisir(request, pk):
    if not _peut_initier(request.user):
        messages.error(request, "Accès refusé.")
        return redirect("inventaires_liste")

    sites = request.user.sites_autorises()
    inv = get_object_or_404(Inventaire.objects.select_related("site"), pk=pk, statut=StatutInventaire.EN_COURS)
    if not sites.filter(pk=inv.site_id).exists():
        messages.error(request, "Accès refusé.")
        return redirect("inventaires_liste")

    if request.method == "POST":
        action = request.POST.get("action", "enregistrer")

        if action == "annuler":
            inv.statut = StatutInventaire.ANNULE
            inv.save(update_fields=["statut"])
            messages.success(request, "Inventaire supprimé.")
            return redirect("inventaires_liste")

        # Sauvegarder les quantités et commentaires dans tous les cas
        for ligne in inv.lignes.all():
            val = request.POST.get(f"qte_{ligne.pk}")
            commentaire = request.POST.get(f"com_{ligne.pk}", "").strip()
            try:
                ligne.quantite_physique = max(0, int(val))
                ligne.commentaire = commentaire
                ligne.save(update_fields=["quantite_physique", "commentaire"])
            except (TypeError, ValueError):
                pass

        if action == "valider":
            # Vérifier que chaque ligne en écart a une raison renseignée
            lignes_sans_raison = [
                l for l in inv.lignes.all()
                if l.ecart != 0 and not l.commentaire.strip()
            ]
            if lignes_sans_raison:
                noms = ", ".join(l.produit.designation for l in lignes_sans_raison[:3])
                messages.error(
                    request,
                    f"Renseignez la raison de l'écart pour : {noms}"
                    + (" (et d'autres)." if len(lignes_sans_raison) > 3 else "."),
                )
                return redirect("inventaire_saisir", pk=pk)
            blocages = verifier_blocages_inventaire(inv.site)
            if blocages:
                for b in blocages:
                    messages.error(request, f"Validation bloquée — {b['message']}.")
                return redirect("inventaire_saisir", pk=pk)
            try:
                valider_inventaire(inv, par=request.user)
                messages.success(request, "Inventaire validé — stock mis à jour.")
            except ValidationError as e:
                messages.error(request, str(e))
            return redirect("inventaires_liste")

        messages.success(request, "Quantités enregistrées.")
        return redirect("inventaire_saisir", pk=pk)

    # Resynchroniser quantite_theorique avec les soldes actuels avant affichage.
    # Un mouvement de stock survenu pendant la saisie (vente, transfert…) modifie
    # SoldeStock mais pas les lignes d'inventaire déjà créées — sans ce refresh
    # la colonne "Théorique" serait périmée.
    soldes_actuels = {
        s.produit_id: s.quantite
        for s in SoldeStock.objects.filter(site=inv.site)
    }
    a_maj = []
    for ligne in inv.lignes.all():
        nouveau = soldes_actuels.get(ligne.produit_id, 0)
        if ligne.quantite_theorique != nouveau:
            ligne.quantite_theorique = nouveau
            a_maj.append(ligne)
    if a_maj:
        InventaireLigne.objects.bulk_update(a_maj, ["quantite_theorique"])

    lignes = inv.lignes.select_related("produit").order_by("produit__code")
    blocages = verifier_blocages_inventaire(inv.site)
    return render(request, "inventaires/saisir.html", {
        "inventaire": inv,
        "lignes": lignes,
        "blocages": blocages,
    })


@login_required
def inventaire_valider(request, pk):
    if not _peut_initier(request.user):
        messages.error(request, "Accès refusé.")
        return redirect("inventaires_liste")

    sites = request.user.sites_autorises()
    inv = get_object_or_404(Inventaire, pk=pk, statut=StatutInventaire.EN_COURS)
    if not sites.filter(pk=inv.site_id).exists():
        messages.error(request, "Accès refusé.")
        return redirect("inventaires_liste")
    if request.method == "POST":
        blocages = verifier_blocages_inventaire(inv.site)
        if blocages:
            for b in blocages:
                messages.error(request, f"Validation bloquée — {b['message']}. Finalisez ce processus avant de valider.")
            return redirect("inventaire_saisir", pk=pk)
        try:
            valider_inventaire(inv, par=request.user)
            messages.success(request, "Inventaire validé — écarts enregistrés au stock.")
        except ValidationError as e:
            messages.error(request, str(e))
        return redirect("inventaires_liste")
    return redirect("inventaire_saisir", pk=pk)


@login_required
def inventaire_detail(request, pk):
    inv = get_object_or_404(
        Inventaire.objects.select_related("site", "cree_par", "valide_par"),
        pk=pk,
    )
    if not request.user.sites_autorises().filter(pk=inv.site_id).exists():
        messages.error(request, "Accès refusé.")
        return redirect("inventaires_liste")
    lignes = inv.lignes.select_related("produit").order_by("produit__code")
    return render(request, "inventaires/detail.html", {
        "inventaire": inv,
        "lignes": lignes,
    })
