"""Exports Excel pour l'approvisionnement."""

import io

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.utils import timezone

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import CommandeEcole, CommandeEcoleLigne, CommandeMagasin, StatutCommandeEcole, StatutCommandeMagasin


def _style_entete(cell):
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.fill = PatternFill("solid", fgColor="16233F")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _ajuster_colonnes(ws):
    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 50)


@login_required
def export_commandes_magasin_excel(request):
    from core.models import Profil

    u = request.user
    sites_autorises = u.sites_autorises()

    qs = (
        CommandeMagasin.objects
        .filter(magasin__in=sites_autorises)
        .select_related("magasin", "cree_par")
        .prefetch_related("lignes__produit")
        .order_by("-cree_le")
    )

    if request.GET.get("magasin", "").isdigit():
        qs = qs.filter(magasin_id=request.GET["magasin"])
    if request.GET.get("statut"):
        qs = qs.filter(statut=request.GET["statut"])
    if request.GET.get("debut"):
        qs = qs.filter(cree_le__date__gte=request.GET["debut"])
    if request.GET.get("fin"):
        qs = qs.filter(cree_le__date__lte=request.GET["fin"])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Commandes magasin"
    ws.freeze_panes = "A2"

    entetes = ["#", "Magasin", "Statut", "Date", "Créé par", "Produit", "Code", "Qté demandée", "Qté livrée", "Qté reçue"]
    for col, titre in enumerate(entetes, 1):
        cell = ws.cell(row=1, column=col, value=titre)
        _style_entete(cell)
    ws.row_dimensions[1].height = 30

    row = 2
    for commande in qs:
        lignes = list(commande.lignes.all())
        if not lignes:
            ws.cell(row=row, column=1, value=commande.pk)
            ws.cell(row=row, column=2, value=commande.magasin.nom)
            ws.cell(row=row, column=3, value=commande.get_statut_display())
            ws.cell(row=row, column=4, value=commande.cree_le.strftime("%d/%m/%Y"))
            ws.cell(row=row, column=5, value=commande.cree_par.get_full_name() or commande.cree_par.username)
            row += 1
        else:
            for ligne in lignes:
                ws.cell(row=row, column=1, value=commande.pk)
                ws.cell(row=row, column=2, value=commande.magasin.nom)
                ws.cell(row=row, column=3, value=commande.get_statut_display())
                ws.cell(row=row, column=4, value=commande.cree_le.strftime("%d/%m/%Y"))
                ws.cell(row=row, column=5, value=commande.cree_par.get_full_name() or commande.cree_par.username)
                ws.cell(row=row, column=6, value=ligne.produit.designation)
                ws.cell(row=row, column=7, value=ligne.produit.code)
                ws.cell(row=row, column=8, value=ligne.quantite)
                ws.cell(row=row, column=9, value=ligne.quantite_livree)
                ws.cell(row=row, column=10, value=ligne.quantite_recue)
                row += 1

    _ajuster_colonnes(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nom = f"commandes_magasin_{timezone.localdate().isoformat()}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nom}"'
    return response


@login_required
def export_commande_magasin_detail_excel(request, pk):
    commande = CommandeMagasin.objects.select_related(
        "magasin", "cree_par"
    ).prefetch_related("lignes__produit").get(pk=pk)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Commande #{commande.numero}"
    ws.freeze_panes = "A2"

    avec_livraison = commande.statut in (StatutCommandeMagasin.LIVREE, StatutCommandeMagasin.RECUE)
    avec_reception = commande.statut == StatutCommandeMagasin.RECUE

    entetes = ["Produit", "Code", "Qté demandée"]
    if avec_livraison:
        entetes.append("Qté livrée")
    if avec_reception:
        entetes.append("Qté reçue")
    for col, titre in enumerate(entetes, 1):
        cell = ws.cell(row=1, column=col, value=titre)
        _style_entete(cell)
    ws.row_dimensions[1].height = 30

    row = 2
    for ligne in commande.lignes.select_related("produit").order_by("produit__code"):
        ws.cell(row=row, column=1, value=ligne.produit.designation)
        ws.cell(row=row, column=2, value=ligne.produit.code)
        ws.cell(row=row, column=3, value=ligne.quantite)
        if avec_livraison:
            ws.cell(row=row, column=4, value=ligne.quantite_livree)
        if avec_reception:
            ws.cell(row=row, column=5, value=ligne.quantite_recue)
        row += 1

    _ajuster_colonnes(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nom = f"commande_magasin_{commande.pk}_{timezone.localdate().isoformat()}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nom}"'
    return response


@login_required
def export_commande_ecole_detail_excel(request, pk):
    commande = CommandeEcole.objects.select_related(
        "ecole", "cree_par"
    ).prefetch_related("lignes__produit").get(pk=pk)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Commande #{commande.numero}"
    ws.freeze_panes = "A2"

    avec_livraison = commande.statut in (StatutCommandeEcole.LIVREE, StatutCommandeEcole.RECUE)
    avec_reception = commande.statut == StatutCommandeEcole.RECUE

    entetes = ["Produit", "Code", "Qté commandée"]
    if avec_livraison:
        entetes.append("Qté livrée")
    if avec_reception:
        entetes.append("Qté reçue")
    for col, titre in enumerate(entetes, 1):
        cell = ws.cell(row=1, column=col, value=titre)
        _style_entete(cell)
    ws.row_dimensions[1].height = 30

    row = 2
    for ligne in commande.lignes.select_related("produit").order_by("produit__code"):
        ws.cell(row=row, column=1, value=ligne.produit.designation)
        ws.cell(row=row, column=2, value=ligne.produit.code)
        ws.cell(row=row, column=3, value=ligne.quantite_demandee)
        if avec_livraison:
            ws.cell(row=row, column=4, value=ligne.quantite_livree)
        if avec_reception:
            ws.cell(row=row, column=5, value=ligne.quantite_recue)
        row += 1

    _ajuster_colonnes(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nom = f"commande_ecole_{commande.pk}_{timezone.localdate().isoformat()}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nom}"'
    return response


@login_required
def export_receptions_magasin_excel(request):
    u = request.user
    sites_autorises = u.sites_autorises()

    qs = (
        CommandeMagasin.objects
        .filter(
            magasin__in=sites_autorises,
            statut__in=[StatutCommandeMagasin.LIVREE, StatutCommandeMagasin.RECUE],
        )
        .select_related("magasin", "livree_par")
        .prefetch_related("lignes__produit")
        .order_by("-livree_le")
    )

    if request.GET.get("magasin", "").isdigit():
        qs = qs.filter(magasin_id=request.GET["magasin"])
    if request.GET.get("statut"):
        qs = qs.filter(statut=request.GET["statut"])
    if request.GET.get("debut"):
        qs = qs.filter(livree_le__date__gte=request.GET["debut"])
    if request.GET.get("fin"):
        qs = qs.filter(livree_le__date__lte=request.GET["fin"])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Réceptions magasin"
    ws.freeze_panes = "A2"

    entetes = ["#", "Magasin", "Statut", "Livrée le", "Livrée par", "Produit", "Code", "Qté livrée", "Qté reçue"]
    for col, titre in enumerate(entetes, 1):
        cell = ws.cell(row=1, column=col, value=titre)
        _style_entete(cell)
    ws.row_dimensions[1].height = 30

    row = 2
    for commande in qs:
        livree_par = commande.livree_par
        livree_par_nom = (livree_par.get_full_name() or livree_par.username) if livree_par else ""
        lignes = list(commande.lignes.all())
        if not lignes:
            ws.cell(row=row, column=1, value=commande.pk)
            ws.cell(row=row, column=2, value=commande.magasin.nom)
            ws.cell(row=row, column=3, value=commande.get_statut_display())
            ws.cell(row=row, column=4, value=commande.livree_le.strftime("%d/%m/%Y") if commande.livree_le else "")
            ws.cell(row=row, column=5, value=livree_par_nom)
            row += 1
        else:
            for ligne in lignes:
                ws.cell(row=row, column=1, value=commande.pk)
                ws.cell(row=row, column=2, value=commande.magasin.nom)
                ws.cell(row=row, column=3, value=commande.get_statut_display())
                ws.cell(row=row, column=4, value=commande.livree_le.strftime("%d/%m/%Y") if commande.livree_le else "")
                ws.cell(row=row, column=5, value=livree_par_nom)
                ws.cell(row=row, column=6, value=ligne.produit.designation)
                ws.cell(row=row, column=7, value=ligne.produit.code)
                ws.cell(row=row, column=8, value=ligne.quantite_livree)
                ws.cell(row=row, column=9, value=ligne.quantite_recue)
                row += 1

    _ajuster_colonnes(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nom = f"receptions_magasin_{timezone.localdate().isoformat()}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nom}"'
    return response


@login_required
def export_livraisons_magasin_excel(request):
    u = request.user
    sites_autorises = u.sites_autorises()

    qs = (
        CommandeMagasin.objects
        .filter(magasin__in=sites_autorises, statut=StatutCommandeMagasin.VALIDEE)
        .select_related("magasin", "validee_par")
        .prefetch_related("lignes__produit")
        .order_by("-cree_le")
    )

    if request.GET.get("magasin", "").isdigit():
        qs = qs.filter(magasin_id=request.GET["magasin"])
    if request.GET.get("debut"):
        qs = qs.filter(cree_le__date__gte=request.GET["debut"])
    if request.GET.get("fin"):
        qs = qs.filter(cree_le__date__lte=request.GET["fin"])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Livraisons magasin"
    ws.freeze_panes = "A2"

    entetes = ["#", "Magasin", "Date création", "Validée par", "Produit", "Code", "Qté demandée"]
    for col, titre in enumerate(entetes, 1):
        cell = ws.cell(row=1, column=col, value=titre)
        _style_entete(cell)
    ws.row_dimensions[1].height = 30

    row = 2
    for commande in qs:
        validee_par = commande.validee_par
        validee_par_nom = (validee_par.get_full_name() or validee_par.username) if validee_par else ""
        lignes = list(commande.lignes.all())
        if not lignes:
            ws.cell(row=row, column=1, value=commande.pk)
            ws.cell(row=row, column=2, value=commande.magasin.nom)
            ws.cell(row=row, column=3, value=commande.cree_le.strftime("%d/%m/%Y"))
            ws.cell(row=row, column=4, value=validee_par_nom)
            row += 1
        else:
            for ligne in lignes:
                ws.cell(row=row, column=1, value=commande.pk)
                ws.cell(row=row, column=2, value=commande.magasin.nom)
                ws.cell(row=row, column=3, value=commande.cree_le.strftime("%d/%m/%Y"))
                ws.cell(row=row, column=4, value=validee_par_nom)
                ws.cell(row=row, column=5, value=ligne.produit.designation)
                ws.cell(row=row, column=6, value=ligne.produit.code)
                ws.cell(row=row, column=7, value=ligne.quantite)
                row += 1

    _ajuster_colonnes(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nom = f"livraisons_magasin_{timezone.localdate().isoformat()}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nom}"'
    return response


@login_required
def export_commandes_ecole_excel(request):
    u = request.user
    sites_autorises = u.sites_autorises()

    qs = (
        CommandeEcole.objects
        .filter(ecole__in=sites_autorises)
        .select_related("ecole", "cree_par")
        .prefetch_related("lignes__produit")
        .order_by("-cree_le")
    )

    if request.GET.get("ecole", "").isdigit():
        qs = qs.filter(ecole_id=request.GET["ecole"])
    if request.GET.get("debut"):
        qs = qs.filter(cree_le__date__gte=request.GET["debut"])
    if request.GET.get("fin"):
        qs = qs.filter(cree_le__date__lte=request.GET["fin"])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Commandes site"
    ws.freeze_panes = "A2"

    entetes = ["#", "Site", "Statut", "Date", "Créé par", "Produit", "Code", "Qté demandée"]
    for col, titre in enumerate(entetes, 1):
        cell = ws.cell(row=1, column=col, value=titre)
        _style_entete(cell)
    ws.row_dimensions[1].height = 30

    row = 2
    for commande in qs:
        lignes = list(commande.lignes.all())
        if not lignes:
            ws.cell(row=row, column=1, value=commande.pk)
            ws.cell(row=row, column=2, value=commande.ecole.nom)
            ws.cell(row=row, column=3, value=commande.get_statut_display())
            ws.cell(row=row, column=4, value=commande.cree_le.strftime("%d/%m/%Y"))
            ws.cell(row=row, column=5, value=commande.cree_par.get_full_name() or commande.cree_par.username)
            row += 1
        else:
            for ligne in lignes:
                ws.cell(row=row, column=1, value=commande.pk)
                ws.cell(row=row, column=2, value=commande.ecole.nom)
                ws.cell(row=row, column=3, value=commande.get_statut_display())
                ws.cell(row=row, column=4, value=commande.cree_le.strftime("%d/%m/%Y"))
                ws.cell(row=row, column=5, value=commande.cree_par.get_full_name() or commande.cree_par.username)
                ws.cell(row=row, column=6, value=ligne.produit.designation)
                ws.cell(row=row, column=7, value=ligne.produit.code)
                ws.cell(row=row, column=8, value=ligne.quantite)
                row += 1

    _ajuster_colonnes(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nom = f"commandes_ecole_{timezone.localdate().isoformat()}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nom}"'
    return response


@login_required
def export_receptions_ecole_excel(request):
    u = request.user
    sites_autorises = u.sites_autorises()

    qs = (
        CommandeEcole.objects
        .filter(
            ecole__in=sites_autorises,
            statut__in=[StatutCommandeEcole.LIVREE, StatutCommandeEcole.RECUE],
        )
        .select_related("ecole", "livree_par")
        .prefetch_related("lignes__produit")
        .order_by("-livree_le")
    )

    if request.GET.get("ecole", "").isdigit():
        qs = qs.filter(ecole_id=request.GET["ecole"])
    if request.GET.get("debut"):
        qs = qs.filter(livree_le__date__gte=request.GET["debut"])
    if request.GET.get("fin"):
        qs = qs.filter(livree_le__date__lte=request.GET["fin"])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Réceptions site"
    ws.freeze_panes = "A2"

    entetes = ["#", "Site", "Statut", "Livrée le", "Livrée par", "Produit", "Code", "Qté livrée"]
    for col, titre in enumerate(entetes, 1):
        cell = ws.cell(row=1, column=col, value=titre)
        _style_entete(cell)
    ws.row_dimensions[1].height = 30

    row = 2
    for commande in qs:
        livree_par = commande.livree_par
        livree_par_nom = (livree_par.get_full_name() or livree_par.username) if livree_par else ""
        lignes = list(commande.lignes.all())
        if not lignes:
            ws.cell(row=row, column=1, value=commande.pk)
            ws.cell(row=row, column=2, value=commande.ecole.nom)
            ws.cell(row=row, column=3, value=commande.get_statut_display())
            ws.cell(row=row, column=4, value=commande.livree_le.strftime("%d/%m/%Y") if commande.livree_le else "")
            ws.cell(row=row, column=5, value=livree_par_nom)
            row += 1
        else:
            for ligne in lignes:
                ws.cell(row=row, column=1, value=commande.pk)
                ws.cell(row=row, column=2, value=commande.ecole.nom)
                ws.cell(row=row, column=3, value=commande.get_statut_display())
                ws.cell(row=row, column=4, value=commande.livree_le.strftime("%d/%m/%Y") if commande.livree_le else "")
                ws.cell(row=row, column=5, value=livree_par_nom)
                ws.cell(row=row, column=6, value=ligne.produit.designation)
                ws.cell(row=row, column=7, value=ligne.produit.code)
                ws.cell(row=row, column=8, value=ligne.quantite)
                row += 1

    _ajuster_colonnes(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nom = f"receptions_ecole_{timezone.localdate().isoformat()}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nom}"'
    return response


@login_required
def export_livraisons_ecole_excel(request):
    u = request.user
    sites_autorises = u.sites_autorises()

    qs = (
        CommandeEcole.objects
        .filter(
            ecole__in=sites_autorises,
            statut=StatutCommandeEcole.VALIDEE,
        )
        .select_related("ecole", "validee_par")
        .prefetch_related("lignes__produit")
        .order_by("-validee_le")
    )

    if request.GET.get("ecole", "").isdigit():
        qs = qs.filter(ecole_id=request.GET["ecole"])
    if request.GET.get("debut"):
        qs = qs.filter(validee_le__date__gte=request.GET["debut"])
    if request.GET.get("fin"):
        qs = qs.filter(validee_le__date__lte=request.GET["fin"])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Livraisons site"
    ws.freeze_panes = "A2"

    entetes = ["#", "Site", "Validée le", "Validée par", "Produit", "Code", "Qté demandée"]
    for col, titre in enumerate(entetes, 1):
        cell = ws.cell(row=1, column=col, value=titre)
        _style_entete(cell)
    ws.row_dimensions[1].height = 30

    row = 2
    for commande in qs:
        validee_par = commande.validee_par
        validee_par_nom = (validee_par.get_full_name() or validee_par.username) if validee_par else ""
        lignes = list(commande.lignes.all())
        if not lignes:
            ws.cell(row=row, column=1, value=commande.pk)
            ws.cell(row=row, column=2, value=commande.ecole.nom)
            ws.cell(row=row, column=3, value=commande.validee_le.strftime("%d/%m/%Y") if commande.validee_le else "")
            ws.cell(row=row, column=4, value=validee_par_nom)
            row += 1
        else:
            for ligne in lignes:
                ws.cell(row=row, column=1, value=commande.pk)
                ws.cell(row=row, column=2, value=commande.ecole.nom)
                ws.cell(row=row, column=3, value=commande.validee_le.strftime("%d/%m/%Y") if commande.validee_le else "")
                ws.cell(row=row, column=4, value=validee_par_nom)
                ws.cell(row=row, column=5, value=ligne.produit.designation)
                ws.cell(row=row, column=6, value=ligne.produit.code)
                ws.cell(row=row, column=7, value=ligne.quantite)
                row += 1

    _ajuster_colonnes(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nom = f"livraisons_ecole_{timezone.localdate().isoformat()}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nom}"'
    return response
