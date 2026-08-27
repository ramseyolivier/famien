import json

from django.contrib import messages
from django.db import transaction
from django.db.models import Count, Q, Sum
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone

from catalogue.models import CategorieProduit, Marque, Produit, PrixEcole
from core.models import Site, TypeSite, Utilisateur
from kits.models import ClasseEcole, Kit, KitLigne
from stock.models import SoldeStock
from stock.services import references_sous_seuil
from ventes.models import StatutVente, Vente

from .decorateurs import admin_requis
from .forms import (
    CategorieProduitForm,
    ClasseEcoleForm,
    KitCreationForm,
    KitLigneFormSet,
    KitPrixForm,
    MarqueForm,
    ProduitForm,
    ReinitialisationMdpForm,
    SiteForm,
    UtilisateurCreationForm,
    UtilisateurModificationForm,
)


# ─── Tableau de bord ────────────────────────────────────────────────────────


@admin_requis
def accueil(request):
    sites = request.user.sites_autorises()
    ventes = Vente.objects.filter(ecole__in=sites).exclude(statut=StatutVente.ANNULEE)
    aujourdhui = ventes.filter(horodatage__date=timezone.localdate())

    contexte = {
        # ── Pilotage du jour ──────────────────────────────────────────────
        "ca_jour": aujourdhui.aggregate(t=Sum("montant_total"))["t"] or 0,
        "nb_ventes_jour": aujourdhui.count(),
        "ca_total": ventes.aggregate(t=Sum("montant_total"))["t"] or 0,
        "par_site": (
            aujourdhui.values("ecole__nom")
            .annotate(ca=Sum("montant_total"), n=Count("id"))
            .order_by("-ca")
        ),
        "alertes": references_sous_seuil(sites=sites).select_related("site", "produit")[:20],
        "valeur_stock": SoldeStock.objects.filter(site__in=sites).aggregate(v=Sum("quantite"))["v"] or 0,
        # ── Référentiel ───────────────────────────────────────────────────
        "nb_produits": Produit.objects.filter(actif=True).count(),
        "nb_marques": Marque.objects.count(),
        "nb_categories": CategorieProduit.objects.count(),
        "nb_kits": Kit.objects.filter(actif=True).count(),
        "nb_sites": Site.objects.filter(actif=True).count(),
        "nb_classes": ClasseEcole.objects.count(),
        "nb_utilisateurs": Utilisateur.objects.filter(is_active=True).count(),
        "derniers_utilisateurs": Utilisateur.objects.filter(is_active=True).order_by("-last_login")[:5],
    }
    return render(request, "administration/accueil.html", contexte)


# ─── Sites ────────────────────────────────────────────────────────────────────


@admin_requis
def sites_liste(request):
    sites = Site.objects.all()
    return render(request, "administration/sites/liste.html", {"sites": sites})


@admin_requis
def site_formulaire(request, pk=None):
    instance = get_object_or_404(Site, pk=pk) if pk else None
    titre = "Modifier le site" if instance else "Nouveau site"
    if request.method == "POST":
        form = SiteForm(request.POST, instance=instance)
        if form.is_valid():
            form.save()
            messages.success(request, "Site enregistré.")
            return redirect("admin_sites")
    else:
        form = SiteForm(instance=instance)
    return render(request, "administration/sites/formulaire.html", {"form": form, "titre": titre, "instance": instance})


@admin_requis
def site_activer(request, pk):
    site = get_object_or_404(Site, pk=pk)
    if request.method == "POST":
        site.actif = not site.actif
        site.save()
        etat = "activé" if site.actif else "désactivé"
        messages.success(request, f"Site {etat}.")
    return redirect("admin_sites")


@admin_requis
def site_tarifs(request, pk):
    from decimal import Decimal, InvalidOperation
    site = get_object_or_404(Site, pk=pk)
    produits = list(Produit.objects.filter(actif=True).select_related("categorie", "marque").order_by("code"))

    if request.method == "POST":
        prix_existants = {pe.produit_id: pe for pe in PrixEcole.objects.filter(ecole=site)}
        a_creer, a_maj = [], []
        a_supprimer = []
        for produit in produits:
            valeur = request.POST.get(f"prix_{produit.pk}", "").strip().replace(" ", "").replace("\xa0", "")
            if valeur:
                try:
                    prix = Decimal(valeur)
                    if prix < 0:
                        continue
                    if produit.pk in prix_existants:
                        pe = prix_existants[produit.pk]
                        pe.prix_detail = prix
                        a_maj.append(pe)
                    else:
                        a_creer.append(PrixEcole(produit=produit, ecole=site, prix_detail=prix))
                except (ValueError, InvalidOperation):
                    pass
            else:
                if produit.pk in prix_existants:
                    a_supprimer.append(prix_existants[produit.pk].pk)
        if a_creer:
            PrixEcole.objects.bulk_create(a_creer)
        if a_maj:
            PrixEcole.objects.bulk_update(a_maj, ["prix_detail"])
        if a_supprimer:
            PrixEcole.objects.filter(pk__in=a_supprimer).delete()
        if request.headers.get("X-Requested-With") == "XMLHttpRequest":
            from django.http import JsonResponse
            return JsonResponse({"ok": True})
        messages.success(request, "Tarifs enregistrés.")
        return redirect("admin_site_tarifs", pk=pk)

    prix_existants = {pe.produit_id: pe.prix_detail for pe in PrixEcole.objects.filter(ecole=site)}
    lignes = [
        {
            "produit": p,
            "prix_ecole": prix_existants.get(p.pk),
            "cout_js": str(int(p.cout_achat)),
            "prix_js": str(int(prix_existants[p.pk])) if p.pk in prix_existants else "",
        }
        for p in produits
    ]
    return render(request, "administration/sites/tarifs.html", {
        "site": site,
        "lignes": lignes,
    })


@admin_requis
def export_site_tarifs_excel(request, pk):
    import io
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    site = get_object_or_404(Site, pk=pk)
    produits = list(Produit.objects.filter(actif=True).select_related("categorie").order_by("categorie__nom", "code"))
    prix_existants = {pe.produit_id: pe.prix_detail for pe in PrixEcole.objects.filter(ecole=site)}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tarifs"
    ws.freeze_panes = "A2"

    entetes = ["Code", "Désignation", "Catégorie", "Prix achat (F CFA)", "Prix vente (F CFA)", "Marge (F CFA)"]
    for col, titre in enumerate(entetes, 1):
        cell = ws.cell(row=1, column=col, value=titre)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor="16233F")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    for row, p in enumerate(produits, 2):
        pv = prix_existants.get(p.pk)
        marge = (pv - p.cout_achat) if pv is not None else None
        ws.cell(row=row, column=1, value=p.code)
        ws.cell(row=row, column=2, value=p.designation)
        ws.cell(row=row, column=3, value=p.categorie.nom if p.categorie_id else "")
        ws.cell(row=row, column=4, value=int(p.cout_achat))
        ws.cell(row=row, column=5, value=int(pv) if pv is not None else "")
        ws.cell(row=row, column=6, value=int(marge) if marge is not None else "")

    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 50)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nom = f"tarifs_{site.nom.lower().replace(' ', '_')}_{timezone.localdate().isoformat()}.xlsx"
    from django.http import HttpResponse
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nom}"'
    return response


# ─── Catégories produit ───────────────────────────────────────────────────────


@admin_requis
def categories_liste(request):
    categories = CategorieProduit.objects.all()
    return render(request, "administration/categories/liste.html", {"categories": categories})


@admin_requis
def categorie_formulaire(request, pk=None):
    instance = get_object_or_404(CategorieProduit, pk=pk) if pk else None
    titre = "Modifier la catégorie" if instance else "Nouvelle catégorie"
    if request.method == "POST":
        form = CategorieProduitForm(request.POST, instance=instance)
        if form.is_valid():
            form.save()
            messages.success(request, "Catégorie enregistrée.")
            return redirect("admin_categories")
    else:
        form = CategorieProduitForm(instance=instance)
    return render(request, "administration/categories/formulaire.html", {"form": form, "titre": titre, "instance": instance})


@admin_requis
def categorie_supprimer(request, pk):
    categorie = get_object_or_404(CategorieProduit, pk=pk)
    if request.method == "POST":
        try:
            categorie.delete()
            messages.success(request, f"Catégorie « {categorie.nom} » supprimée.")
        except Exception:
            messages.error(request, "Impossible de supprimer cette catégorie : des produits l'utilisent encore.")
        return redirect("admin_categories")
    return render(request, "administration/categories/liste.html", {
        "categories": CategorieProduit.objects.all(),
        "supprimer": categorie,
    })


# ─── Marques ─────────────────────────────────────────────────────────────────


@admin_requis
def marques_liste(request):
    marques = Marque.objects.all()
    return render(request, "administration/marques/liste.html", {"marques": marques})


@admin_requis
def marque_formulaire(request, pk=None):
    instance = get_object_or_404(Marque, pk=pk) if pk else None
    titre = "Modifier la marque" if instance else "Nouvelle marque"
    if request.method == "POST":
        form = MarqueForm(request.POST, instance=instance)
        if form.is_valid():
            form.save()
            messages.success(request, "Marque enregistrée.")
            return redirect("admin_marques")
    else:
        form = MarqueForm(instance=instance)
    return render(request, "administration/marques/formulaire.html", {"form": form, "titre": titre, "instance": instance})


@admin_requis
def marque_supprimer(request, pk):
    marque = get_object_or_404(Marque, pk=pk)
    if request.method == "POST":
        try:
            marque.delete()
            messages.success(request, f"Marque « {marque.nom} » supprimée.")
        except Exception:
            messages.error(request, "Impossible de supprimer cette marque : des produits l'utilisent encore.")
        return redirect("admin_marques")
    return render(request, "administration/marques/liste.html", {
        "marques": Marque.objects.all(),
        "supprimer": marque,
    })


# ─── Produits ────────────────────────────────────────────────────────────────


@admin_requis
def produits_liste(request):
    produits = Produit.objects.select_related("marque", "categorie").order_by("categorie__nom", "code")
    return render(request, "administration/produits/liste.html", {"produits": produits})


@admin_requis
def produit_formulaire(request, pk=None):
    instance = get_object_or_404(Produit, pk=pk) if pk else None
    titre = "Modifier le produit" if instance else "Nouveau produit"
    if request.method == "POST":
        form = ProduitForm(request.POST, instance=instance)
        if form.is_valid():
            form.save()
            messages.success(request, "Produit enregistré.")
            return redirect("admin_produits")
    else:
        form = ProduitForm(instance=instance)
    return render(request, "administration/produits/formulaire.html", {"form": form, "titre": titre, "instance": instance})


@admin_requis
def produit_activer(request, pk):
    produit = get_object_or_404(Produit, pk=pk)
    if request.method == "POST":
        produit.actif = not produit.actif
        produit.save()
        etat = "activé" if produit.actif else "désactivé"
        messages.success(request, f"Produit {etat}.")
    return redirect("admin_produits")


@admin_requis
def produit_tarifs(request, pk):
    from decimal import Decimal, InvalidOperation
    produit = get_object_or_404(Produit, pk=pk)
    sites = list(Site.objects.filter(actif=True).order_by("nom"))

    if request.method == "POST":
        prix_existants = {pe.ecole_id: pe for pe in PrixEcole.objects.filter(produit=produit)}
        a_creer, a_maj, a_supprimer = [], [], []
        for site in sites:
            valeur = request.POST.get(f"prix_{site.pk}", "").strip().replace(" ", "").replace("\xa0", "")
            if valeur:
                try:
                    prix = Decimal(valeur)
                    if prix < 0:
                        continue
                    if site.pk in prix_existants:
                        pe = prix_existants[site.pk]
                        pe.prix_detail = prix
                        a_maj.append(pe)
                    else:
                        a_creer.append(PrixEcole(produit=produit, ecole=site, prix_detail=prix))
                except (ValueError, InvalidOperation):
                    pass
            else:
                if site.pk in prix_existants:
                    a_supprimer.append(prix_existants[site.pk].pk)
        if a_creer:
            PrixEcole.objects.bulk_create(a_creer)
        if a_maj:
            PrixEcole.objects.bulk_update(a_maj, ["prix_detail"])
        if a_supprimer:
            PrixEcole.objects.filter(pk__in=a_supprimer).delete()
        messages.success(request, "Tarifs enregistrés.")
        return redirect("admin_produit_tarifs", pk=pk)

    prix_existants = {pe.ecole_id: pe.prix_detail for pe in PrixEcole.objects.filter(produit=produit)}
    lignes = [
        {
            "site": s,
            "prix_ecole": prix_existants.get(s.pk),
        }
        for s in sites
    ]
    return render(request, "administration/produits/tarifs.html", {
        "produit": produit,
        "lignes": lignes,
    })


# ─── Classes d'école ─────────────────────────────────────────────────────────


@admin_requis
def classes_liste(request):
    classes = ClasseEcole.objects.select_related("ecole").prefetch_related("kits").order_by("ecole__nom", "niveau", "libelle")
    return render(request, "administration/kits/classes_liste.html", {"classes": classes})


@admin_requis
def classe_formulaire(request, pk=None):
    instance = get_object_or_404(ClasseEcole, pk=pk) if pk else None
    titre = "Modifier la classe" if instance else "Nouvelle classe"
    if request.method == "POST":
        form = ClasseEcoleForm(request.POST, instance=instance)
        if form.is_valid():
            form.save()
            messages.success(request, "Classe enregistrée.")
            return redirect("admin_classes")
    else:
        form = ClasseEcoleForm(instance=instance)
    return render(request, "administration/kits/classe_formulaire.html", {"form": form, "titre": titre, "instance": instance})


@admin_requis
def classe_supprimer(request, pk):
    classe = get_object_or_404(ClasseEcole, pk=pk)
    if request.method == "POST":
        if classe.kits.exists():
            messages.error(request, "Impossible de supprimer une classe qui possède des kits.")
        else:
            classe.delete()
            messages.success(request, "Classe supprimée.")
        return redirect("admin_classes")
    return redirect("admin_classes")


# ─── Kits ────────────────────────────────────────────────────────────────────


@admin_requis
def kits_liste(request):
    kits = Kit.objects.filter(actif=True).select_related("ecole", "classe").prefetch_related("lignes__produit")
    return render(request, "administration/kits/liste.html", {"kits": kits})


@admin_requis
def kit_formulaire(request, pk=None):
    kit_existant = get_object_or_404(Kit, pk=pk, actif=True) if pk else None
    titre = "Modifier le kit" if kit_existant else "Nouveau kit"

    if request.method == "POST":
        formset = KitLigneFormSet(request.POST, prefix="ligne")
        if kit_existant:
            prix_form = KitPrixForm(request.POST)
            creation_form = None
            formulaires_ok = prix_form.is_valid() and formset.is_valid()
        else:
            creation_form = KitCreationForm(request.POST)
            prix_form = None
            formulaires_ok = creation_form.is_valid() and formset.is_valid()

        if formulaires_ok:
            lignes_valides = [
                f.cleaned_data for f in formset
                if f.cleaned_data and not f.cleaned_data.get("DELETE") and f.cleaned_data.get("produit")
            ]
            if not lignes_valides:
                messages.error(request, "Un kit doit contenir au moins un article.")
            else:
                try:
                    with transaction.atomic():
                        if kit_existant:
                            nouveau = kit_existant.nouvelle_version(auteur=request.user)
                            nouveau.prix_vente = prix_form.cleaned_data["prix_vente"]
                            nouveau.save()
                            nouveau.lignes.all().delete()
                        else:
                            classe = creation_form.cleaned_data["classe"]
                            nouveau = Kit(
                                ecole=classe.ecole,
                                classe=classe,
                                prix_vente=creation_form.cleaned_data["prix_vente"],
                                cree_par=request.user,
                            )
                            nouveau.full_clean()
                            nouveau.save()
                        KitLigne.objects.bulk_create([
                            KitLigne(kit=nouveau, produit=l["produit"], quantite=l["quantite"])
                            for l in lignes_valides
                        ])
                        nouveau.verifier_regle_de_prix()
                    messages.success(request, "Kit enregistré.")
                    return redirect("admin_kits")
                except Exception as e:
                    messages.error(request, str(e))
    else:
        formset = KitLigneFormSet(prefix="ligne")
        if kit_existant:
            prix_form = KitPrixForm(initial={"prix_vente": kit_existant.prix_vente})
            creation_form = None
            initial_lignes = [
                {"produit": l.produit, "quantite": l.quantite}
                for l in kit_existant.lignes.all()
            ]
            formset = KitLigneFormSet(initial=initial_lignes, prefix="ligne")
        else:
            creation_form = KitCreationForm()
            prix_form = None

    return render(request, "administration/kits/formulaire.html", {
        "creation_form": creation_form,
        "prix_form": prix_form,
        "formset": formset,
        "titre": titre,
        "kit_existant": kit_existant,
        "produits_json": _produits_json(),
        "classes_json": _classes_json(),
    })


def _classes_json():
    """Classes par ecole_id — pour filtrage JS dans le formulaire de création de kit."""
    data = {}
    for c in ClasseEcole.objects.select_related("ecole").order_by("niveau", "libelle"):
        data.setdefault(str(c.ecole_id), []).append({"id": c.pk, "libelle": str(c)})
    return json.dumps(data)


def _produits_json():
    data = {
        str(p.pk): {"cout": float(p.cout_achat), "prix_detail": float(p.prix_detail), "prix_ecoles": {}}
        for p in Produit.objects.filter(actif=True)
    }
    for pe in PrixEcole.objects.select_related("produit").filter(produit__actif=True):
        key = str(pe.produit_id)
        if key in data:
            data[key]["prix_ecoles"][str(pe.ecole_id)] = float(pe.prix_detail)
    return json.dumps(data)


@admin_requis
def kit_supprimer(request, pk):
    kit = get_object_or_404(Kit, pk=pk, actif=True)
    if request.method == "POST":
        Kit.objects.filter(pk=pk).update(actif=False)
        messages.success(request, "Kit désactivé.")
        return redirect("admin_kits")
    return redirect("admin_kits")


# ─── Utilisateurs ────────────────────────────────────────────────────────────


@admin_requis
def utilisateurs_connexions(request):
    from core.models import Profil
    profil_filtre = request.GET.get("profil", "")
    q = request.GET.get("q", "").strip()
    tri = request.GET.get("tri", "desc")
    ordre = "derniere_activite" if tri == "asc" else "-derniere_activite"
    utilisateurs = Utilisateur.objects.filter(is_active=True).order_by(ordre).select_related("site")
    if profil_filtre:
        utilisateurs = utilisateurs.filter(profil=profil_filtre)
    if q:
        utilisateurs = utilisateurs.filter(
            Q(first_name__icontains=q) | Q(last_name__icontains=q) | Q(username__icontains=q)
        )
    return render(request, "administration/utilisateurs/connexions.html", {
        "utilisateurs": utilisateurs,
        "profils": Profil.choices,
        "filtres": {"profil": profil_filtre, "q": q, "tri": tri},
    })


@admin_requis
def utilisateurs_liste(request):
    utilisateurs = Utilisateur.objects.select_related("site").order_by("profil", "last_name", "first_name")
    return render(request, "administration/utilisateurs/liste.html", {"utilisateurs": utilisateurs})


@admin_requis
def utilisateur_formulaire(request, pk=None):
    instance = get_object_or_404(Utilisateur, pk=pk) if pk else None
    titre = "Modifier l'utilisateur" if instance else "Nouvel utilisateur"
    FormClass = UtilisateurModificationForm if instance else UtilisateurCreationForm
    if request.method == "POST":
        form = FormClass(request.POST, instance=instance)
        if form.is_valid():
            form.save()
            messages.success(request, "Utilisateur enregistré.")
            return redirect("admin_utilisateurs")
    else:
        form = FormClass(instance=instance)
    return render(request, "administration/utilisateurs/formulaire.html", {
        "form": form,
        "titre": titre,
        "instance": instance,
        "sites_json": _sites_json(),
    })


def _sites_json():
    """Sites actifs pour le formulaire utilisateur."""
    data = [
        {"id": s.pk, "nom": str(s)}
        for s in Site.objects.filter(actif=True).order_by("nom")
    ]
    return json.dumps(data)


@admin_requis
def utilisateur_activer(request, pk):
    utilisateur = get_object_or_404(Utilisateur, pk=pk)
    if request.method == "POST":
        utilisateur.is_active = not utilisateur.is_active
        utilisateur.save()
        etat = "activé" if utilisateur.is_active else "désactivé"
        messages.success(request, f"Compte {etat}.")
    return redirect("admin_utilisateurs")


@admin_requis
def utilisateur_reset_mdp(request, pk):
    utilisateur = get_object_or_404(Utilisateur, pk=pk)
    if request.method == "POST":
        form = ReinitialisationMdpForm(request.POST)
        if form.is_valid():
            utilisateur.set_password(form.cleaned_data["nouveau_mot_de_passe"])
            utilisateur.save()
            messages.success(request, "Mot de passe réinitialisé.")
            return redirect("admin_utilisateurs")
    else:
        form = ReinitialisationMdpForm()
    return render(request, "administration/utilisateurs/reset_mdp.html", {
        "form": form,
        "utilisateur": utilisateur,
    })
