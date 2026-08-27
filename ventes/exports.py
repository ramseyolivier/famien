"""Exports Excel pour les rapports de ventes (M24)."""

import io
from datetime import date as date_cls

from django.contrib.auth.decorators import login_required
from django.db.models import Sum
from django.http import HttpResponse
from django.utils import timezone

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.models import Profil
from kits.models import Niveau
from .models import StatutVente, Vente


def _style_entete(cell):
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.fill = PatternFill("solid", fgColor="16233F")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _ajuster_colonnes(ws):
    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 50)


@login_required
def export_ventes_excel(request):
    from ventes.views import _ecoles_perimetre

    ecoles = _ecoles_perimetre(request.user)
    qs = (
        Vente.objects
        .filter(ecole__in=ecoles)
        .select_related("ecole", "vendeuse")
        .order_by("-horodatage")
    )

    debut = request.GET.get("debut", "")
    fin = request.GET.get("fin", "")
    statut = request.GET.get("statut", "")
    if debut:
        try:
            qs = qs.filter(horodatage__date__gte=date_cls.fromisoformat(debut))
        except ValueError:
            pass
    if fin:
        try:
            qs = qs.filter(horodatage__date__lte=date_cls.fromisoformat(fin))
        except ValueError:
            pass
    if statut:
        qs = qs.filter(statut=statut)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ventes"
    ws.freeze_panes = "A2"

    entetes = ["N° Vente", "Site", "Vendeuse", "Date", "Heure", "Mode paiement", "Montant", "Remise", "Statut"]
    for col, titre in enumerate(entetes, 1):
        cell = ws.cell(row=1, column=col, value=titre)
        _style_entete(cell)
    ws.row_dimensions[1].height = 30

    for row, v in enumerate(qs[:5000], 2):
        ws.cell(row=row, column=1, value=v.numero)
        ws.cell(row=row, column=2, value=v.ecole.nom)
        ws.cell(row=row, column=3, value=v.vendeuse.get_full_name() or v.vendeuse.username)
        ws.cell(row=row, column=4, value=v.horodatage.date())
        ws.cell(row=row, column=4).number_format = "DD/MM/YYYY"
        ws.cell(row=row, column=5, value=timezone.localtime(v.horodatage).strftime("%H:%M"))
        ws.cell(row=row, column=6, value=v.get_mode_paiement_display())
        ws.cell(row=row, column=7, value=int(v.montant_total))
        ws.cell(row=row, column=8, value=int(v.remise) if v.remise else 0)
        ws.cell(row=row, column=9, value=v.get_statut_display())

    _ajuster_colonnes(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nom = f"ventes_{timezone.localdate().isoformat()}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nom}"'
    return response


@login_required
def export_kits_excel(request):
    from ventes.views import _ecoles_perimetre, _kits_disponibles
    from stock.models import SoldeStock
    from core.models import Profil

    u = request.user
    toutes_ecoles = _ecoles_perimetre(u).filter(actif=True).order_by("nom")

    ecole_id = request.GET.get("ecole", "")
    niveau_filtre = request.GET.get("niveau", "")
    non_constructibles = request.GET.get("nc", "")

    ecoles = toutes_ecoles
    if ecole_id:
        ecoles = ecoles.filter(pk=ecole_id)

    magasin_stock = None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kits constructibles"
    ws.freeze_panes = "A2"

    entetes = ["Site", "Niveau", "Prix vente (F CFA)", "Constructibles", "Articles manquants"]
    for col, titre in enumerate(entetes, 1):
        cell = ws.cell(row=1, column=col, value=titre)
        _style_entete(cell)
    ws.row_dimensions[1].height = 30

    rouge = PatternFill("solid", fgColor="FEE2E2")
    row = 2
    for ecole in ecoles:
        kits = _kits_disponibles(ecole, stock_depuis=magasin_stock)
        if niveau_filtre:
            kits = [k for k in kits if k["kit"].niveau == niveau_filtre]
        if non_constructibles:
            kits = [k for k in kits if k["constructibles"] == 0]
        for k in kits:
            ws.cell(row=row, column=1, value=ecole.nom)
            ws.cell(row=row, column=2, value=k["kit"].classe.libelle)
            ws.cell(row=row, column=3, value=int(k["kit"].prix_vente))
            cell_c = ws.cell(row=row, column=4, value=k["constructibles"])
            if k["constructibles"] == 0:
                cell_c.fill = rouge
            ws.cell(row=row, column=5, value=", ".join(k["manquants"]) if k["manquants"] else "")
            row += 1

    _ajuster_colonnes(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nom = f"kits_{timezone.localdate().isoformat()}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nom}"'
    return response


@login_required
def export_annulations_excel(request):
    from ventes.views import _ecoles_perimetre

    ecoles = _ecoles_perimetre(request.user)
    qs = (
        Vente.objects
        .filter(statut=StatutVente.ANNULEE, ecole__in=ecoles)
        .select_related("ecole", "vendeuse", "annulee_par")
        .prefetch_related("lignes")
        .order_by("-annulee_le")
    )

    debut      = request.GET.get("debut", "")
    fin        = request.GET.get("fin", "")
    ecole_id   = request.GET.get("ecole", "")

    if debut:
        qs = qs.filter(annulee_le__date__gte=debut)
    if fin:
        qs = qs.filter(annulee_le__date__lte=fin)
    if ecole_id.isdigit():
        qs = qs.filter(ecole_id=ecole_id)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Annulations"
    ws.freeze_panes = "A2"

    entetes = ["Annulée le", "N° Vente", "Site", "Caissière", "Annulée par", "Montant", "Motif"]
    for col, titre in enumerate(entetes, 1):
        cell = ws.cell(row=1, column=col, value=titre)
        _style_entete(cell)
    ws.row_dimensions[1].height = 30

    for row, v in enumerate(qs[:5000], 2):
        ws.cell(row=row, column=1, value=timezone.localtime(v.annulee_le).strftime("%d/%m/%Y %H:%M") if v.annulee_le else "")
        ws.cell(row=row, column=2, value=v.numero)
        ws.cell(row=row, column=3, value=v.ecole.nom)
        ws.cell(row=row, column=4, value=v.vendeuse.get_full_name() or v.vendeuse.username)
        ws.cell(row=row, column=5, value=v.annulee_par.get_full_name() if v.annulee_par else "")
        ws.cell(row=row, column=6, value=int(v.montant_total))
        ws.cell(row=row, column=7, value=v.motif_annulation)

    _ajuster_colonnes(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nom = f"annulations_{timezone.localdate().isoformat()}.xlsx"
    response = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nom}"'
    return response
